import copy
from io import BytesIO
from types import SimpleNamespace as NS
from docx import Document
import pytest
from app.services.document_generator import generate_document, generation_preview


def fixture():
    cat=NS(id='c',sku='SW8',brand='FABRICANTE TESTE',model='SW8',category='switch',description='',specification='{"atributos":{"portas_rj45":{"valor":8}}}',lpu_version=None)
    state={'categoria':'switch','requisitos':[{'id':'R1','atributo':'portas_rj45','operador':'>=','valor':24,'texto_original':'No mínimo 24 portas RJ45','fonte':{'pagina':12}}], 'filtros':[], 'descricao_revisada':True,'obrigacoes_gerais':{'acessorios':'Kit rack'},'obrigacoes_revisadas':False}
    item=NS(id='i',item_number='11',description='Switch',raw_payload={'_manual_match':state},catalog_product=cat)
    return NS(id='n',number='PE 189/2026',bid_number='189/2026',organ=NS(name='Órgão teste'),notice_products=[item],municipality_name='São Paulo'), {'item_ids':['i']}


def test_preview_restricts_to_one_known_saved_model_and_never_uses_other_item():
    notice,opts=fixture()
    assert generation_preview(notice,'complementary_declaration',{}, {'item_ids':['foreign']})['missing_fields']==['declaration.saved_match']
    assert generation_preview(notice,'complementary_declaration',{}, {'item_ids':['i','i']})['missing_fields']==['declaration.item_selection']
    notice.notice_products[0]._declaration_catalog=None
    assert generation_preview(notice,'complementary_declaration',{},opts)['missing_fields']==['declaration.saved_match']


def test_draft_preserves_literal_requirement_failure_identity_and_match_state():
    notice,opts=fixture();before=copy.deepcopy(notice.notice_products[0].raw_payload)
    content,preview=generate_document(notice,'complementary_declaration',{},opts)
    assert preview['fields']['declaration']['points'][0]['status']=='NAO_ATENDE'
    assert notice.notice_products[0].raw_payload==before
    doc=Document(BytesIO(content));text='\n'.join(p.text for p in doc.paragraphs)+'\n'+'\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    for required in ['MINUTA PARA REVISÃO DA TOR','TOR TECNOLOGIA','No mínimo 24 portas RJ45','Não declarar atendimento','FABRICANTE TESTE','SKU: SW8','[A CONFIRMAR]','Kit rack','pagina: 12']:
        assert required in text
    assert '510 W' not in text and 'RAISECOM DO BRASIL' not in text
    assert all(len(list(section.header._element.iter('{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}anchor')))==1 for section in doc.sections)


def test_generator_rejects_missing_selection_even_when_allow_missing_requested():
    notice,_=fixture()
    with pytest.raises(ValueError):generate_document(notice,'complementary_declaration',{}, {'allow_missing':True})
