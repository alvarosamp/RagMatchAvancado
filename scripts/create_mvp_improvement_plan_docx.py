from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "Plano_Melhoria_MVP_CRM_IA_Sistema.docx"

BLACK = "000000"
NAVY = "17365D"
MID_BLUE = "2F5597"
LIGHT_BLUE = "EAF1F8"
PALE_BLUE = "F5F8FC"
LIGHT_GRAY = "F2F2F2"
MID_GRAY = "666666"
BORDER = "D9D9D9"
WHITE = "FFFFFF"
GREEN = "E2F0D9"
AMBER = "FFF2CC"
RED = "FCE4D6"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_borders(cell, color: str = BORDER, size: str = "6") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn("w:" + margin))
        if node is None:
            node = OxmlElement("w:" + margin)
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_col_width(cell, width_cm: float) -> None:
    cell.width = Cm(width_cm)
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(int(width_cm * 567)))
    tc_w.set(qn("w:type"), "dxa")


def format_run(run, *, bold=None, italic=None, size=None, color=None, font="Arial") -> None:
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), font)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), font)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if size is not None:
        run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    format_run(run, size=8.5, color=MID_GRAY)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def remove_paragraph_borders(paragraph_or_style) -> None:
    element = paragraph_or_style._element
    p_pr = element.find(qn("w:pPr"))
    if p_pr is None:
        return
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is not None:
        p_pr.remove(p_bdr)


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.1)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor.from_string(BLACK)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    title = styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(27)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(BLACK)
    title.paragraph_format.space_after = Pt(14)
    remove_paragraph_borders(title)

    for style_name, size, before, after in (
        ("Heading 1", 17, 18, 8),
        ("Heading 2", 13.5, 14, 6),
        ("Heading 3", 11.5, 10, 4),
    ):
        st = styles[style_name]
        st.font.name = "Arial"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = RGBColor.from_string(BLACK)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True

    if "Small Note" not in styles:
        small = styles.add_style("Small Note", WD_STYLE_TYPE.PARAGRAPH)
    else:
        small = styles["Small Note"]
    small.font.name = "Arial"
    small._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    small._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    small.font.size = Pt(8.5)
    small.font.color.rgb = RGBColor.from_string(MID_GRAY)
    small.paragraph_format.space_after = Pt(5)

    for sec in doc.sections:
        footer = sec.footer
        add_page_number(footer.paragraphs[0])


def add_heading(doc: Document, text: str, level: int = 1) -> None:
    doc.add_heading(text, level=level)


def add_para(doc: Document, text: str = "", *, bold_lead: str | None = None, style=None, align=None) -> None:
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        format_run(r1, bold=True)
        r2 = p.add_run(text[len(bold_lead):])
        format_run(r2)
    else:
        r = p.add_run(text)
        format_run(r)


def add_bullets(doc: Document, items: list[str], *, level: int = 0) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.paragraph_format.space_after = Pt(3.5)
        p.paragraph_format.left_indent = Cm(0.65 + level * 0.45)
        p.paragraph_format.first_line_indent = Cm(-0.25)
        r = p.add_run(item)
        format_run(r, size=10.2)


def add_numbered(doc: Document, items: list[str]) -> None:
    for index, item in enumerate(items, start=1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.75)
        p.paragraph_format.first_line_indent = Cm(-0.55)
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(f"{index}.  {item}")
        format_run(r, size=10.2)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None,
              header_fill: str = NAVY, font_size: float = 8.8) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.style = "Table Grid"
    hdr = table.rows[0]
    repeat_table_header(hdr)
    for idx, header in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_shading(cell, header_fill)
        set_cell_borders(cell)
        set_cell_margins(cell, 120, 120, 120, 120)
        if widths:
            set_col_width(cell, widths[idx])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        format_run(r, bold=True, size=9, color=WHITE)
    prevent_row_split(hdr)

    for ridx, row_data in enumerate(rows):
        row = table.add_row()
        prevent_row_split(row)
        fill = WHITE if ridx % 2 == 0 else PALE_BLUE
        for idx, value in enumerate(row_data):
            cell = row.cells[idx]
            set_cell_shading(cell, fill)
            set_cell_borders(cell)
            set_cell_margins(cell)
            if widths:
                set_col_width(cell, widths[idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx == 0 and len(headers) > 2 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(str(value))
            format_run(r, size=font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_checklist(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.35)
        p.paragraph_format.first_line_indent = Cm(-0.35)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run("☐  " + item)
        format_run(r, size=9.8, font="Arial")


def page_break(doc: Document) -> None:
    doc.add_page_break()


def build_document() -> Document:
    doc = Document()
    style_document(doc)

    # Cover
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    remove_paragraph_borders(p)
    r = p.add_run("Plano de melhoria e execução do MVP")
    format_run(r, bold=True, size=27, color=BLACK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(20)
    r = p.add_run("CRM IA Sistema Arquitetura e estratégia de produto")
    format_run(r, size=15, color=MID_GRAY)

    add_para(doc, "Documento para discussão entre os dois desenvolvedores responsáveis pelo produto", style="Small Note")
    add_para(doc, f"Versão de trabalho preparada em {date.today().strftime('%d/%m/%Y')}", style="Small Note")

    doc.add_paragraph().paragraph_format.space_after = Pt(28)
    add_para(
        doc,
        "Este documento organiza as decisões necessárias para transformar a plataforma atual em um MVP comercial. "
        "A recomendação central é concentrar o produto no fluxo de licitações, manter uma arquitetura multiempresa e "
        "configurável e adiar módulos que não sejam necessários para provar valor com clientes reais.",
    )
    add_para(
        doc,
        "Ao final da reunião, os dois desenvolvedores devem sair com um público inicial definido, um escopo congelado, "
        "responsáveis por cada frente, métricas de sucesso e uma data para iniciar o piloto.",
    )

    page_break(doc)

    # Contents and meeting agenda
    add_heading(doc, "Sumário", 1)
    contents = [
        "1 Resumo executivo",
        "2 Diagnóstico do produto atual",
        "3 Posicionamento e mercado inicial",
        "4 Definição do MVP",
        "5 Plano da frente CRM",
        "6 Plano da frente IA",
        "7 Plano da frente Sistema",
        "8 System design e arquitetura",
        "9 Dados multiempresa e integrações",
        "10 Segurança privacidade e conformidade",
        "11 Observabilidade desempenho e operação",
        "12 Qualidade e estratégia de testes",
        "13 Roadmap de execução",
        "14 Organização do trabalho em dupla",
        "15 Métricas do produto e critérios de lançamento",
        "16 Estratégia de piloto e comercialização",
        "17 Riscos e respostas",
        "18 Backlog priorizado",
        "19 Pauta da reunião de alinhamento",
        "20 Registro de decisões",
        "21 Próximos passos em 30 60 e 90 dias",
        "22 Glossário e referências internas",
    ]
    add_bullets(doc, contents)

    add_heading(doc, "Decisões que precisam sair da reunião", 2)
    add_checklist(doc, [
        "Quem é o primeiro perfil de cliente e qual problema ele compra.",
        "Qual é o único fluxo que precisa funcionar de ponta a ponta no MVP.",
        "Quais funcionalidades ficam explicitamente fora da primeira versão.",
        "Qual desenvolvedor lidera cada frente e onde haverá revisão conjunta.",
        "Quais métricas autorizam o início do piloto.",
        "Qual é a data alvo do primeiro uso com dados de cliente.",
    ])

    page_break(doc)

    # 1
    add_heading(doc, "1 Resumo executivo", 1)
    add_para(
        doc,
        "A plataforma já possui uma base técnica maior do que a de um MVP inicial. O repositório contém um backend FastAPI, "
        "frontend React, PostgreSQL com pgvector, Redis, MinIO, filas Dramatiq, autenticação multiempresa, CRM de editais, "
        "processamento de documentos, matching híbrido e recursos de MLOps. O trabalho prioritário não é acrescentar tecnologia. "
        "É reduzir o escopo, fechar a jornada principal e tornar o produto seguro para cadastrar uma nova empresa sem alteração de código.",
    )
    add_para(
        doc,
        "O produto inicial deve atender empresas que participam de licitações públicas. A arquitetura pode suportar outros mercados, "
        "mas tentar vender o MVP para qualquer tipo de empresa levaria o time a construir um CRM genérico, um gerenciador documental "
        "genérico e uma plataforma de IA genérica ao mesmo tempo. Esse escopo não é adequado para dois desenvolvedores.",
    )
    add_para(doc, "Recomendação de produto", bold_lead="Recomendação de produto")
    add_para(
        doc,
        "Entregar uma plataforma que encontra ou recebe editais, extrai os requisitos, cruza os itens com o catálogo do cliente, "
        "apresenta evidências para revisão e conduz a oportunidade pelo CRM até o resultado final.",
    )

    add_heading(doc, "Resultado esperado do MVP", 2)
    add_table(
        doc,
        ["Dimensão", "Resultado necessário"],
        [
            ["Produto", "Um fluxo principal compreensível e demonstrável em menos de 15 minutos."],
            ["CRM", "Uma oportunidade percorre todo o funil com prazos, itens, documentos, histórico e resultado."],
            ["IA", "A extração e o matching mostram evidências, confiança e opção de correção humana."],
            ["Sistema", "Uma nova empresa é ativada por configuração, com isolamento de dados e operação observável."],
            ["Comercial", "Duas ou três empresas piloto usam dados reais e confirmam economia de tempo."],
        ],
        widths=[3.1, 12.2],
        font_size=9.2,
    )

    add_heading(doc, "Princípios de decisão", 2)
    add_numbered(doc, [
        "Cada nova funcionalidade deve contribuir diretamente para captar, analisar, decidir ou acompanhar uma oportunidade.",
        "A IA deve apoiar uma decisão verificável. Respostas sem evidência não são entregáveis do produto.",
        "O sistema deve funcionar sem LLM, ainda que com menos automação.",
        "Configuração por empresa deve substituir valores fixos, mas customização ilimitada fica fora do MVP.",
        "O monólito modular deve ser mantido enquanto o volume e a organização do time não justificarem microserviços.",
        "O piloto deve começar antes de novos módulos avançados entrarem no backlog.",
    ])

    # 2
    add_heading(doc, "2 Diagnóstico do produto atual", 1)
    add_para(
        doc,
        "A análise do repositório em 15 de setembro de 2026 identificou 106 arquivos Python no backend, 63 arquivos JavaScript ou JSX "
        "no frontend principal e 39 arquivos de testes. Esse volume confirma que a plataforma está em fase de consolidação, não de criação "
        "de uma prova de conceito.",
    )

    add_heading(doc, "Capacidades existentes", 2)
    add_bullets(doc, [
        "Autenticação com tenant e papéis de administrador, editor e visualizador.",
        "CRM com editais, órgãos, portais, produtos, documentos, histórico, sessões, resultados e pós disputa.",
        "Upload e processamento de PDFs, extração estruturada e armazenamento de análises.",
        "Matching por regras, embeddings, pgvector e LLM configurável.",
        "Filas assíncronas, worker dedicado de IA e recuperação de jobs.",
        "Armazenamento de arquivos em MinIO ou S3 e exportações.",
        "PgBouncer, Redis, Traefik, health checks e imagens de produção.",
        "Feature flags para ativar IA, extração por LLM e matching por LLM.",
        "Métricas de matching, avaliação e monitoramento de drift com MLflow e Evidently.",
    ])

    add_heading(doc, "Lacunas com impacto comercial", 2)
    add_table(
        doc,
        ["Lacuna", "Impacto", "Tratamento recomendado"],
        [
            ["Dados empresariais ainda fixos em partes do gerador", "Impede ativação limpa de novos clientes", "Criar perfil empresarial por tenant e eliminar valores fixos"],
            ["Amplitude elevada de módulos", "Dificulta explicar e concluir o MVP", "Ocultar ou pausar módulos fora da jornada principal"],
            ["Fluxos ainda aguardam validação com catálogo real", "Risco de falha na demonstração", "Testar vínculo simples, kits, LPU ausente e mínimo manual"],
            ["Atributos técnicos duplicados", "Reduz confiança e leitura operacional", "Adotar chaves canônicas, prioridade de fontes e tratamento de conflito"],
            ["IA desabilitada por padrão em produção", "Valor central ainda não está provado em operação", "Reativar por tenant depois de criar avaliação e limites"],
            ["Configuração de mercado parcialmente codificada", "Generalização vira manutenção de código", "Mover configurações essenciais para tenant e feature flags"],
            ["Tarefas e alertas ainda incompletos", "Prazos podem depender de controle externo", "Incluir apenas tarefa simples e alertas críticos no P1"],
        ],
        widths=[4.3, 4.5, 6.6],
        font_size=8.3,
    )

    add_heading(doc, "Conclusão do diagnóstico", 2)
    add_para(
        doc,
        "O repositório já contém os blocos necessários. A melhoria deve priorizar coerência do domínio, isolamento multiempresa, "
        "qualidade mensurável da IA e uma operação simples. Uma reescrita ampla ou migração para microserviços consumiria tempo sem resolver "
        "a principal incerteza, que é a adesão do cliente ao fluxo proposto.",
    )

    page_break(doc)

    # 3
    add_heading(doc, "3 Posicionamento e mercado inicial", 1)
    add_heading(doc, "Mercado recomendado", 2)
    add_para(
        doc,
        "O primeiro mercado deve ser formado por empresas privadas que vendem produtos ou serviços para órgãos públicos e mantêm um "
        "catálogo técnico. O melhor perfil inicial é uma equipe pequena ou média que analisa vários editais, compara requisitos manualmente "
        "e perde tempo consolidando informações em planilhas, e-mails e pastas.",
    )

    add_heading(doc, "Perfil de cliente inicial", 2)
    add_table(
        doc,
        ["Critério", "Perfil desejável"],
        [
            ["Atividade", "Participa de pregões e licitações com frequência"],
            ["Catálogo", "Possui produtos, serviços ou soluções que podem ser descritos por atributos"],
            ["Equipe", "Entre 2 e 20 pessoas envolvidas em análise, documentação ou disputa"],
            ["Dor", "Gasta horas lendo documentos e comparando especificações"],
            ["Dados", "Consegue fornecer editais anteriores, catálogo e decisões corretas para avaliação"],
            ["Patrocinador", "Existe um responsável disposto a testar semanalmente e dar feedback"],
        ],
        widths=[3.6, 11.7],
        font_size=9.1,
    )

    add_heading(doc, "Promessa comercial inicial", 2)
    add_para(
        doc,
        "A plataforma reduz o tempo entre encontrar um edital e decidir como tratá-lo. Ela centraliza o documento, extrai os pontos críticos, "
        "compara os itens com o catálogo e mantém a equipe organizada até o encerramento.",
    )

    add_heading(doc, "Limite da generalização", 2)
    add_para(
        doc,
        "O núcleo técnico deve aceitar múltiplas empresas, provedores de IA e perfis de mercado. Entretanto, os nomes, regras e telas do MVP "
        "devem continuar orientados a licitações. Um perfil genérico de matching documental pode permanecer como direção futura, sem condicionar "
        "o lançamento atual.",
    )

    add_heading(doc, "Hipóteses que o piloto deve validar", 2)
    add_checklist(doc, [
        "O cliente considera a extração automática suficientemente confiável para iniciar a análise.",
        "O ranking de produtos reduz o número de comparações manuais.",
        "As evidências permitem revisar a recomendação sem reler o edital inteiro.",
        "O CRM substitui pelo menos parte da planilha ou controle paralelo.",
        "O ganho de tempo justifica pagamento recorrente.",
    ])

    # 4
    add_heading(doc, "4 Definição do MVP", 1)
    add_heading(doc, "Jornada principal", 2)
    architecture = [
        "1  Empresa cadastra usuários e dados empresariais",
        "2  Empresa importa catálogo e datasheets",
        "3  Usuário encontra ou envia um edital",
        "4  Sistema cria um job e processa o documento",
        "5  IA extrai itens prazos documentos e riscos",
        "6  Motor cruza requisitos com o catálogo",
        "7  Usuário revisa evidências e corrige recomendações",
        "8  Oportunidade entra no pipeline do CRM",
        "9  Equipe acompanha preparação disputa e homologação",
        "10 Resultado é registrado como ganho perdido ou cancelado",
    ]
    for line in architecture:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.space_after = Pt(3)
        r = p.add_run(line)
        format_run(r, font="Courier New", size=9.6)

    add_heading(doc, "Escopo obrigatório", 2)
    add_table(
        doc,
        ["Área", "Incluído no MVP"],
        [
            ["Empresa", "Cadastro, usuários, papéis, dados empresariais e preferências básicas"],
            ["Catálogo", "Cadastro, importação, preço, atributos técnicos e datasheet"],
            ["Entrada", "Upload de PDF e captação de oportunidade por fonte já integrada"],
            ["Análise", "Extração de campos críticos, itens, prazos, documentos e riscos"],
            ["Matching", "Ranking de produtos, score, evidências, incerteza e revisão humana"],
            ["CRM", "Pipeline, oportunidade, itens, documentos, responsável, histórico e resultado"],
            ["Operação", "Jobs, reprocessamento, logs, métricas, backup e alertas essenciais"],
        ],
        widths=[3.2, 12.1],
        font_size=9,
    )

    add_heading(doc, "Escopo posterior", 2)
    add_bullets(doc, [
        "Chat avançado com múltiplos documentos.",
        "Inteligência competitiva e previsão de vitória.",
        "Robô autônomo de lances.",
        "Pós venda completo com faturamento e logística.",
        "Construtor visual de pipelines e campos.",
        "Integrações profundas com ERP e portais privados.",
        "White label completo.",
        "Fine tuning e modelos exclusivos por cliente.",
    ])

    add_heading(doc, "Regra para aceitar novas funcionalidades", 2)
    add_para(
        doc,
        "Uma funcionalidade só entra no MVP se resolver uma falha do fluxo principal, for necessária para um piloto confirmado ou reduzir um "
        "risco técnico que pode comprometer dados, segurança ou operação. Pedidos cosméticos e módulos novos ficam no backlog até o piloto produzir evidência.",
    )

    # 5 CRM
    add_heading(doc, "5 Plano da frente CRM", 1)
    add_para(
        doc,
        "O CRM deve ser a área operacional do produto. Ele recebe o resultado da análise e organiza o trabalho humano. O usuário não deveria "
        "precisar copiar dados para uma planilha para continuar o processo.",
    )

    add_heading(doc, "Objetivos do CRM", 2)
    add_bullets(doc, [
        "Mostrar em segundos quais oportunidades exigem ação.",
        "Preservar a identidade do item original do edital.",
        "Permitir vincular um ou mais produtos do catálogo a cada item.",
        "Controlar prazos, documentos, responsáveis e decisões.",
        "Registrar uma trilha compreensível do que aconteceu.",
        "Separar o resultado de cada item do encerramento geral do edital.",
    ])

    add_heading(doc, "Funcionalidades prioritárias", 2)
    add_table(
        doc,
        ["Prioridade", "Capacidade", "Critério de aceite"],
        [
            ["P0", "Pipeline e calendário", "Mesmas oportunidades, filtros no servidor e abertura rápida"],
            ["P0", "Detalhe da oportunidade", "Informações, itens, documentos, histórico e resultado acessíveis sem inconsistência"],
            ["P0", "Itens e catálogo", "Vínculo simples e kit preservam quantidade, lote, referência e identidade do item"],
            ["P0", "Preços", "LPU, mínimo operacional, referência e totais usam nomes e cálculos consistentes"],
            ["P0", "Fluxo e encerramento", "Avanço de fase, descarte, reabertura e resultado respeitam as regras definidas"],
            ["P0", "Dados empresariais", "Documentos gerados usam apenas dados do tenant"],
            ["P1", "Tarefas", "Responsável, prazo, status e vínculo com oportunidade"],
            ["P1", "Alertas", "Pregão próximo, prazo vencido, documento pendente e item sem produto"],
            ["P1", "Dashboard", "Indicadores operacionais e financeiros com filtros simples"],
            ["P2", "Automação comercial", "Regras configuráveis depois de validar tarefas e alertas"],
        ],
        widths=[1.6, 4.2, 9.6],
        font_size=8.3,
    )

    add_heading(doc, "Regras de negócio que precisam ser congeladas", 2)
    add_numbered(doc, [
        "O item importado do edital permanece a entidade comercial principal.",
        "Dois ou mais produtos vinculados ao mesmo item formam um kit e não criam novas linhas do edital.",
        "O mínimo unitário segue a prioridade manual, soma das LPUs do kit e valor de referência.",
        "Um kit com componente sem LPU não pode usar soma parcial sem aviso explícito.",
        "Resultado por item não encerra automaticamente o edital.",
        "Ganho, perda ou cancelamento geral ocorre por ação explícita e fica no histórico.",
        "Descarte é uma decisão interna de triagem. Cancelamento representa encerramento externo.",
    ])

    add_heading(doc, "Validações com dados reais", 2)
    add_checklist(doc, [
        "Oportunidade sem produto vinculado.",
        "Item com um produto vinculado.",
        "Kit com dois ou mais produtos e LPUs diferentes.",
        "Kit com componente sem preço.",
        "Preço mínimo informado manualmente.",
        "Remoção e substituição de componente.",
        "Item retirado e lote totalmente retirado.",
        "Edital descartado, reaberto, encerrado e reaberto após resultado.",
        "Dois usuários editando ou consultando a mesma oportunidade.",
    ])

    add_heading(doc, "Experiência do usuário", 2)
    add_para(
        doc,
        "A página inicial do CRM deve responder três perguntas: o que vence primeiro, o que está bloqueado e quem precisa agir. O cartão do pipeline "
        "deve mostrar apenas informação necessária para priorizar. Informações completas ficam no detalhe e as abas devem carregar sob demanda.",
    )

    # 6 IA
    add_heading(doc, "6 Plano da frente IA", 1)
    add_para(
        doc,
        "A frente de IA deve ser tratada como um sistema de decisão assistida. O valor está na extração, comparação e evidência. Chat, resumo e geração "
        "de texto são complementos posteriores.",
    )

    add_heading(doc, "Pipeline recomendado", 2)
    add_table(
        doc,
        ["Etapa", "Entrada", "Saída", "Falha tratável"],
        [
            ["Ingestão", "PDF ou documento", "Arquivo versionado e hash", "Arquivo inválido ou duplicado"],
            ["Parsing", "Arquivo", "Texto estruturado por página e seção", "OCR parcial ou layout não reconhecido"],
            ["Normalização", "Texto", "Campos canônicos e itens", "Campo ausente ou conflito"],
            ["Embeddings", "Trechos e catálogo", "Vetores com versão", "Modelo indisponível"],
            ["Matching rápido", "Requisitos e produtos", "Candidatos e score", "Poucos candidatos"],
            ["Enriquecimento LLM", "Casos ambíguos", "Justificativa estruturada", "Timeout ou resposta inválida"],
            ["Validação", "Resultado", "Evidências e confiança", "Ausência de fonte"],
            ["Revisão humana", "Recomendação", "Rótulo correto e correção", "Usuário discorda"],
        ],
        widths=[2.5, 3.5, 5.0, 4.4],
        font_size=8.1,
    )

    add_heading(doc, "Contrato mínimo da recomendação", 2)
    add_bullets(doc, [
        "Identificador do item e do produto candidato.",
        "Score final e componentes do score.",
        "Requisitos atendidos, não atendidos e não verificáveis.",
        "Trecho de evidência e origem para cada conclusão crítica.",
        "Nível de confiança.",
        "Versão do parser, embeddings, prompt e modelo.",
        "Tempo, custo estimado e status do processamento.",
        "Decisão humana e motivo da correção, quando houver.",
    ])

    add_heading(doc, "Estratégia de qualidade", 2)
    add_numbered(doc, [
        "Criar um conjunto de referência com editais, itens, produtos corretos e justificativas aprovadas.",
        "Separar exemplos de desenvolvimento e validação para não medir no mesmo material usado no ajuste.",
        "Medir extração por campo e matching por item, em vez de usar apenas uma nota média geral.",
        "Registrar falsos positivos, falsos negativos e casos sem evidência.",
        "Bloquear promoção de prompt ou modelo quando houver regressão acima do limite acordado.",
        "Comparar modelos pelo ganho de qualidade, tempo e custo total por edital.",
    ])

    add_heading(doc, "Metas iniciais sugeridas", 2)
    add_table(
        doc,
        ["Métrica", "Meta inicial", "Observação"],
        [
            ["Campos críticos extraídos corretamente", "95 por cento ou mais", "Datas, órgão, identificadores, itens e quantidades"],
            ["Produto correto no Top 3", "85 por cento ou mais", "Medido no conjunto revisado"],
            ["Conclusões técnicas com evidência", "100 por cento", "Sem fonte, o resultado fica como não verificável"],
            ["Jobs com estado final observável", "95 por cento ou mais", "Sucesso ou falha clara e reprocessável"],
            ["Casos ambíguos enviados ao LLM", "Somente quando necessário", "Controlar custo e fila"],
            ["Decisões críticas automáticas", "Zero", "Revisão humana permanece obrigatória no MVP"],
        ],
        widths=[5.1, 3.4, 6.9],
        font_size=8.6,
    )

    add_heading(doc, "Governança da IA", 2)
    add_bullets(doc, [
        "Feature flag por tenant e por capacidade.",
        "Limite de concorrência, timeout e retry com backoff.",
        "Schema estruturado e validação antes de persistir.",
        "Prompt e modelo versionados.",
        "Nenhum documento de outro tenant em contexto, cache ou busca vetorial.",
        "Logs sem conteúdo integral sensível.",
        "Fallback determinístico quando o provedor estiver indisponível.",
        "Orçamento e consumo acompanhados por tenant.",
    ])

    add_heading(doc, "Funcionalidades adiadas", 2)
    add_para(
        doc,
        "Fine tuning, agentes autônomos, previsão de vitória, geração automática de decisões e robô de lances devem ficar fora do MVP. "
        "Essas capacidades exigem dados históricos, governança e responsabilidade operacional que o piloto ainda precisa produzir.",
    )

    # 7 System
    add_heading(doc, "7 Plano da frente Sistema", 1)
    add_para(
        doc,
        "A frente Sistema reúne as capacidades que permitem vender e operar o produto para mais de uma empresa. Ela não deve virar uma coleção de telas administrativas. "
        "Seu objetivo é tornar o onboarding, a segurança, a configuração e o suporte previsíveis.",
    )

    add_heading(doc, "Capacidades obrigatórias", 2)
    add_table(
        doc,
        ["Capacidade", "Requisito do MVP"],
        [
            ["Onboarding", "Criar tenant, administrador, dados empresariais e configuração inicial sem editar código"],
            ["Usuários", "Convidar, desativar e alterar papel com auditoria"],
            ["Configuração", "Catálogo, checklists, provedor de IA, limites e funcionalidades por tenant"],
            ["Jobs", "Status, progresso, tentativas, erro legível e reprocessamento"],
            ["Arquivos", "Upload privado, versionamento mínimo e download autorizado"],
            ["Auditoria", "Autor, data e alteração em ações críticas"],
            ["Suporte", "Identificadores de correlação e diagnóstico sem acessar conteúdo desnecessário"],
            ["Operação", "Health checks, métricas, alertas, backup e restauração testada"],
        ],
        widths=[3.4, 11.9],
        font_size=9,
    )

    add_heading(doc, "Configuração por empresa", 2)
    add_bullets(doc, [
        "Razão social, nome fantasia, CNPJ, endereço e contatos.",
        "Responsáveis e assinantes autorizados.",
        "Logotipo e cores básicas, sem white label completo.",
        "Catálogo, unidades, preços e regras de mínimo.",
        "Checklists e modelos de documentos permitidos.",
        "Módulos habilitados e limites de uso.",
        "Provedor e modelo de IA autorizados.",
        "Retenção de arquivos e preferências de notificação.",
    ])

    add_heading(doc, "Administração interna", 2)
    add_para(
        doc,
        "O time precisa de uma visão interna simples para ativar um tenant, consultar saúde de jobs, bloquear acesso, revisar consumo e identificar falhas. "
        "Essa área deve evitar acesso casual aos documentos do cliente e registrar qualquer ação administrativa sensível.",
    )

    # 8 Architecture
    add_heading(doc, "8 System design e arquitetura", 1)
    add_heading(doc, "Arquitetura recomendada para o MVP", 2)
    add_para(
        doc,
        "A arquitetura atual deve evoluir como monólito modular. O frontend, a API e os workers podem ser implantados separadamente, mas compartilham um modelo de domínio e uma base de código. "
        "Essa abordagem reduz custo operacional e preserva limites claros entre os módulos.",
    )

    arch_lines = [
        "Usuário",
        "  |",
        "  v",
        "CDN WAF TLS e Traefik",
        "  |",
        "  +--> Frontend React",
        "  |",
        "  +--> API FastAPI stateless",
        "          |",
        "          +--> PostgreSQL e pgvector",
        "          +--> Redis para cache filas e rate limit",
        "          +--> MinIO ou S3 para documentos",
        "          +--> Jobs Dramatiq",
        "                  +--> Worker de documentos",
        "                  +--> Worker de matching",
        "                  +--> Worker de IA",
        "                  +--> Worker de exportação",
        "                          |",
        "                          +--> Gateway interno de IA",
        "                                  +--> Ollama local",
        "                                  +--> Provedor externo futuro",
    ]
    for line in arch_lines:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.45)
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(line)
        format_run(r, font="Courier New", size=9.2)

    add_heading(doc, "Responsabilidade dos componentes", 2)
    add_table(
        doc,
        ["Componente", "Responsabilidade", "Regra de escala"],
        [
            ["Frontend", "Experiência, cache de tela e atualização otimista", "CDN e assets versionados"],
            ["API", "Autorização, validação, casos de uso e resposta rápida", "Réplicas stateless após medir p95"],
            ["PostgreSQL", "Fonte de verdade transacional", "Índices e queries antes de aumentar máquina"],
            ["pgvector", "Busca semântica isolada por tenant", "Manter no Postgres enquanto atender SLO"],
            ["Redis", "Fila, cache e limites", "Descartável para cache e persistente para fila conforme política"],
            ["MinIO ou S3", "Arquivos privados e exportações", "URL assinada e política de retenção"],
            ["Workers", "Processamento assíncrono idempotente", "Escalar por profundidade e tempo de fila"],
            ["Gateway de IA", "Provedor, credenciais, orçamento, timeout e auditoria", "Adicionar provedores sem mudar o domínio"],
            ["MLflow", "Versão, experimentos e métricas", "Não usar como fonte de verdade do CRM"],
        ],
        widths=[3.2, 6.2, 5.9],
        font_size=8.2,
    )

    add_heading(doc, "Limites dos módulos", 2)
    add_bullets(doc, [
        "Identity controla tenants, usuários, autenticação e papéis.",
        "CRM controla oportunidades, fases, responsáveis, histórico e resultados.",
        "Catalog controla produtos, preços, atributos e datasheets.",
        "Documents controla arquivos, versões, parsing e evidências.",
        "Matching controla candidatos, scores, explicações, revisão e avaliação.",
        "Jobs controla execução assíncrona, progresso, tentativas e falhas.",
        "Reports lê dados consolidados sem alterar regras transacionais.",
        "Integrations encapsula PNCP, e-mail, portais e sistemas externos.",
    ])

    add_heading(doc, "Princípios técnicos", 2)
    add_numbered(doc, [
        "Nenhum parsing, embedding ou LLM deve bloquear uma requisição HTTP.",
        "Lista e detalhe devem ter contratos diferentes. A lista retorna resumos e o detalhe carrega relacionamentos sob demanda.",
        "Toda escrita crítica deve ser idempotente ou possuir uma chave de idempotência.",
        "Banco é a fonte de verdade. Cache pode ser descartado e reconstruído.",
        "API e workers devem permanecer stateless sempre que possível.",
        "Integrações externas devem ser acessadas por adapters com timeout, retry e circuit breaker simples.",
        "Mudanças de banco devem seguir expandir, migrar e contrair para reduzir indisponibilidade.",
    ])

    add_heading(doc, "Decisões que não devem ser tomadas agora", 2)
    add_bullets(doc, [
        "Separar cada módulo em um microserviço.",
        "Adicionar Kafka ou outra plataforma de eventos distribuída.",
        "Adotar um banco vetorial separado sem evidência de limite do pgvector.",
        "Criar infraestrutura Kubernetes antes de existir necessidade operacional.",
        "Manter sessão presa a uma instância da API.",
        "Executar LLM dentro do processo web.",
    ])

    # 9 Data
    add_heading(doc, "9 Dados multiempresa e integrações", 1)
    add_heading(doc, "Modelo multiempresa", 2)
    add_para(
        doc,
        "Cada registro de negócio deve possuir tenant_id obrigatório. O filtro por tenant não pode depender da disciplina de cada desenvolvedor em cada query. "
        "A evolução recomendada é centralizar escopo de tenant em repositórios ou casos de uso e adicionar testes automáticos de isolamento.",
    )

    add_heading(doc, "Regras de isolamento", 2)
    add_checklist(doc, [
        "Toda tabela de negócio possui tenant_id e índice adequado.",
        "Toda relação valida que pai e filho pertencem ao mesmo tenant.",
        "Toda chave Redis contém tenant e parâmetros relevantes.",
        "Todo objeto no bucket usa prefixo do tenant.",
        "Todo job carrega tenant_id e valida o recurso antes de processar.",
        "Toda busca vetorial filtra tenant antes de ordenar similaridade.",
        "Toda exportação e URL assinada verifica autorização no momento da emissão.",
        "Logs evitam conteúdo integral e identificadores sensíveis desnecessários.",
    ])

    add_heading(doc, "Entidades mínimas", 2)
    add_table(
        doc,
        ["Domínio", "Entidades principais"],
        [
            ["Identidade", "Tenant, User, RoleAudit"],
            ["Empresa", "CompanyProfile, Signer, TenantSetting, FeatureFlag"],
            ["Catálogo", "CatalogProduct, ProductAttribute, ProductPrice, ProductDatasheet"],
            ["Oportunidade", "Notice, NoticeItem, StageTransition, Assignment, Note"],
            ["Documentos", "Document, DocumentVersion, DocumentChunk, Requirement, Evidence"],
            ["Matching", "MatchRun, MatchCandidate, MatchEvidence, HumanReview"],
            ["Operação", "Job, JobAttempt, AuditEvent, Notification"],
            ["Resultado", "ItemResult, NoticeOutcome, Fulfillment posterior"],
        ],
        widths=[3.5, 11.8],
        font_size=8.9,
    )

    add_heading(doc, "Integrações", 2)
    add_para(
        doc,
        "Cada integração deve possuir credenciais por ambiente, adapter próprio, timeout, retry limitado e registro de sincronização. O domínio não deve depender do formato bruto de PNCP, e-mail ou portal. "
        "O payload original pode ser preservado para auditoria, enquanto uma camada de normalização produz entidades canônicas.",
    )

    add_heading(doc, "Eventos internos sugeridos", 2)
    add_table(
        doc,
        ["Evento", "Consumidores possíveis"],
        [
            ["DocumentUploaded", "Parsing e antivírus"],
            ["DocumentParsed", "Normalização e embeddings"],
            ["NoticeAnalyzed", "CRM e matching"],
            ["CatalogChanged", "Reindexação de embeddings e invalidação de cache"],
            ["MatchCompleted", "Notificação e tela de revisão"],
            ["MatchReviewed", "Dataset de avaliação e métricas"],
            ["NoticeStageChanged", "Histórico, tarefas e alertas"],
            ["NoticeClosed", "Dashboard e relatório"],
        ],
        widths=[5.0, 10.3],
        font_size=8.9,
    )
    add_para(
        doc,
        "No MVP esses eventos podem ser funções ou mensagens Dramatiq no mesmo repositório. O importante é definir contrato, idempotência e rastreabilidade, não adotar uma nova infraestrutura.",
    )

    # 10 Security
    add_heading(doc, "10 Segurança privacidade e conformidade", 1)
    add_para(
        doc,
        "Documentos de licitação, catálogo, preços e estratégias comerciais são dados sensíveis para o cliente. Segurança precisa fazer parte do critério de lançamento, mesmo antes de certificações formais.",
    )

    add_heading(doc, "Controles mínimos", 2)
    add_table(
        doc,
        ["Tema", "Controle mínimo"],
        [
            ["Autenticação", "Cookie Secure HttpOnly SameSite e expiração adequada"],
            ["Autorização", "RBAC e verificação de tenant em todos os recursos"],
            ["Senha", "Hash forte, política mínima e bloqueio de usuário inativo"],
            ["Uploads", "Limite de tamanho, validação de tipo, nome seguro e varredura quando possível"],
            ["Arquivos", "Bucket privado, URL assinada curta e auditoria de download"],
            ["Segredos", "Variáveis do orquestrador ou cofre, nunca repositório ou imagem"],
            ["Logs", "Sem senha, token, documento integral ou prompt sensível"],
            ["API", "Rate limit por usuário e tenant, CORS restrito e headers de segurança"],
            ["Backup", "Criptografia, retenção, controle de acesso e teste de restauração"],
            ["IA externa", "Política clara sobre envio de dados e opção de desativação por tenant"],
        ],
        widths=[3.4, 11.9],
        font_size=8.7,
    )

    add_heading(doc, "LGPD e governança", 2)
    add_bullets(doc, [
        "Definir finalidade e base legal para dados pessoais de usuários e contatos.",
        "Coletar apenas dados necessários para a operação.",
        "Publicar política de privacidade e canal de atendimento.",
        "Definir retenção e exclusão de contas, logs e documentos.",
        "Registrar fornecedores que processam dados, incluindo provedores de IA.",
        "Preparar exportação e exclusão de dados do tenant mediante processo controlado.",
        "Documentar incidentes e responsabilidades antes do piloto pago.",
    ])

    # 11 Observability
    add_heading(doc, "11 Observabilidade desempenho e operação", 1)
    add_heading(doc, "Indicadores técnicos", 2)
    add_table(
        doc,
        ["Camada", "Indicadores"],
        [
            ["Frontend", "LCP, erros de JavaScript, tempo até conteúdo útil e falhas de API"],
            ["API", "Requests, erros, p50, p95, p99, tamanho de payload e tempo por rota"],
            ["Banco", "Queries lentas, conexões, locks, cache hit e crescimento"],
            ["Fila", "Backlog, idade do job mais antigo, tentativas e falhas"],
            ["IA", "Tempo, tokens, custo, erro de schema, timeout e qualidade"],
            ["Arquivos", "Uso de armazenamento, upload, download e falhas"],
            ["Negócio", "Editais analisados, matches revisados, decisões e tempo economizado"],
        ],
        widths=[3.2, 12.1],
        font_size=8.9,
    )

    add_heading(doc, "Objetivos de serviço iniciais", 2)
    add_bullets(doc, [
        "Lista do CRM com p95 de servidor abaixo de 800 milissegundos em base representativa.",
        "Conteúdo principal do frontend em até 2,5 segundos em conexão comum.",
        "Operações transacionais comuns com p95 abaixo de 500 milissegundos, excluindo jobs.",
        "Noventa e cinco por cento dos jobs terminam em sucesso ou falha explicável e reprocessável.",
        "Nenhum upload ou job pesado bloqueia workers da API.",
        "Alertas para indisponibilidade, crescimento de fila, erro repetido e pouco espaço em disco.",
    ])

    add_heading(doc, "Backup e recuperação", 2)
    add_checklist(doc, [
        "Backup automático do PostgreSQL com retenção definida.",
        "Plano de recuperação pontual quando o ambiente justificar.",
        "Versionamento ou cópia dos objetos críticos.",
        "Restauração testada em ambiente separado.",
        "RPO e RTO acordados entre os desenvolvedores.",
        "Procedimento documentado para perda de banco, arquivos, Redis e servidor.",
    ])

    # 12 Testing
    add_heading(doc, "12 Qualidade e estratégia de testes", 1)
    add_heading(doc, "Pirâmide de testes", 2)
    add_table(
        doc,
        ["Nível", "Cobertura desejada"],
        [
            ["Unidade", "Regras de preço, kits, fases, permissões, parsing e composição de score"],
            ["Integração", "Banco, tenant, filas, object storage, pgvector e providers"],
            ["Contrato", "Schemas entre frontend, API, CRM incorporado e workers"],
            ["Ponta a ponta", "Cadastro, upload, análise, revisão, CRM e encerramento"],
            ["Carga", "Lista do CRM, detalhe, upload concorrente e backlog de jobs"],
            ["Segurança", "Acesso cruzado entre tenants, papéis, arquivos e URLs assinadas"],
            ["IA", "Dataset versionado, regressão por campo e ranking"],
        ],
        widths=[3.0, 12.3],
        font_size=8.9,
    )

    add_heading(doc, "Cenários obrigatórios antes do piloto", 2)
    add_checklist(doc, [
        "Usuário do tenant A não acessa recurso, arquivo, cache ou job do tenant B.",
        "Upload duplicado não cria processamento inconsistente.",
        "Job reiniciado não duplica itens ou histórico.",
        "LLM indisponível não derruba o CRM.",
        "Resposta inválida do LLM não é persistida como decisão válida.",
        "Migração de banco preserva dados existentes.",
        "Kit e mínimo operacional produzem o mesmo resultado no backend, frontend e exportação.",
        "Encerramento e reabertura preservam o histórico.",
        "Restauração do backup produz um ambiente utilizável.",
    ])

    page_break(doc)

    # 13 Roadmap
    add_heading(doc, "13 Roadmap de execução", 1)
    add_para(
        doc,
        "O cronograma abaixo pressupõe dedicação regular dos dois desenvolvedores, manutenção limitada durante o ciclo e acesso rápido a empresas piloto. "
        "Se essas condições não existirem, preservem a sequência e ajustem as datas, sem executar todas as fases em paralelo.",
    )

    add_table(
        doc,
        ["Período", "Objetivo", "Entregas", "Critério de saída"],
        [
            ["Semana 1", "Congelar o MVP", "Público, jornada, backlog, métricas e pilotos", "Documento de escopo aprovado pelos dois"],
            ["Semanas 2 e 3", "Completar a base SaaS", "Perfil empresarial, onboarding, tenant, RBAC, auditoria e backup", "Nova empresa ativada sem código"],
            ["Semanas 3 e 4", "Fechar o CRM", "Pipeline, itens, kits, preços, documentos e resultados", "Fluxo completo com dados reais"],
            ["Semanas 4 a 6", "Tornar a IA mensurável", "Dataset, extração, matching, evidência, confiança e feedback", "Metas mínimas atingidas no conjunto de avaliação"],
            ["Semana 7", "Integrar e endurecer", "E2E, carga, segurança, observabilidade e runbooks", "Release candidate sem bloqueadores P0"],
            ["Semana 8", "Iniciar piloto", "Onboarding de duas ou três empresas e acompanhamento", "Primeiro ciclo real concluído"],
            ["Semanas 9 e 10", "Corrigir com evidência", "Ajustes P0, métricas de adoção e proposta comercial", "Decisão de cobrar, iterar ou reposicionar"],
        ],
        widths=[2.2, 3.3, 5.5, 4.4],
        font_size=7.9,
    )

    add_heading(doc, "Marcos de controle", 2)
    add_numbered(doc, [
        "Marco um: qualquer tenant novo completa o onboarding sem alteração de código.",
        "Marco dois: um edital real percorre análise, revisão e CRM até o resultado.",
        "Marco três: a IA atinge as metas mínimas no dataset versionado.",
        "Marco quatro: testes de isolamento, backup e carga são aprovados.",
        "Marco cinco: duas empresas usam o produto e relatam valor mensurável.",
    ])

    add_heading(doc, "Política de mudanças durante o ciclo", 2)
    add_para(
        doc,
        "Pedidos novos entram em uma fila de descoberta. Somente erros que bloqueiam o piloto, riscos de dados ou exigências de um piloto confirmado interrompem o plano. "
        "Toda exceção deve registrar motivo, impacto e item que será retirado para manter a capacidade do time.",
    )

    # 14 Team
    add_heading(doc, "14 Organização do trabalho em dupla", 1)
    add_heading(doc, "Divisão sugerida", 2)
    add_table(
        doc,
        ["Frente", "Liderança inicial", "Responsabilidades"],
        [
            ["CRM e experiência", "Desenvolvedor A", "Fluxo, frontend, regras de oportunidade, catálogo e validação com usuário"],
            ["IA e plataforma", "Desenvolvedor B", "Parsing, matching, avaliação, jobs, observabilidade e infraestrutura"],
            ["Decisões compartilhadas", "Ambos", "Modelo de dados, contratos, tenant, segurança, deploy e escopo"],
        ],
        widths=[3.5, 3.5, 8.3],
        font_size=8.9,
    )

    add_heading(doc, "Cronograma dos dois desenvolvedores", 2)
    add_para(
        doc,
        "O cronograma distribui liderança sem criar silos. Cada desenvolvedor possui uma entrega principal por semana, revisa a mudança crítica do outro e participa de um resultado integrado. "
        "As semanas representam sequência de execução, não datas fixas. Se a disponibilidade for parcial, mantenham a ordem e ampliem a duração de cada etapa.",
    )
    add_table(
        doc,
        ["Semana", "Desenvolvedor A", "Desenvolvedor B", "Entrega conjunta e critério de saída"],
        [
            [
                "1",
                "Mapear a jornada do CRM, telas essenciais, fases, itens, kits, preços e regras de encerramento.",
                "Mapear arquitetura, modelo de dados, isolamento por tenant, pipeline de IA e riscos operacionais.",
                "Escopo P0, contratos principais, métricas e empresas piloto aprovados pelos dois.",
            ],
            [
                "2",
                "Implementar perfil empresarial, onboarding, preferências e fluxo inicial de catálogo no frontend.",
                "Concluir tenant, RBAC, auditoria, migrações, configuração e remoção dos dados empresariais fixos.",
                "Uma nova empresa é ativada, configura seus dados e acessa apenas os próprios recursos.",
            ],
            [
                "3",
                "Fechar lista e detalhe do CRM, filtros, fases, responsável, histórico e estados de interface.",
                "Estabilizar upload, jobs, parsing, normalização, idempotência, retries e diagnóstico de falhas.",
                "Um edital enviado aparece no CRM com status rastreável e pode ser reprocessado com segurança.",
            ],
            [
                "4",
                "Concluir vínculo entre item e catálogo, kits, LPU, preço mínimo, descarte e reabertura.",
                "Implementar matching base, busca vetorial, composição de score e armazenamento de evidências.",
                "Um item real recebe candidatos do catálogo sem perder identidade, quantidade ou preço.",
            ],
            [
                "5",
                "Construir revisão humana de recomendações, correção, justificativa, confiança e estados vazios ou de erro.",
                "Criar dataset versionado e avaliação de extração, Top 3, falsos positivos, falsos negativos e evidência.",
                "Relatório de qualidade reproduzível e revisão humana persistida no histórico.",
            ],
            [
                "6",
                "Finalizar documentos, tarefas e alertas P1 indispensáveis, resultado por item e encerramento geral.",
                "Concluir gateway de IA, fallback determinístico, feature flags, limites, timeout e custo por tenant.",
                "Fluxo completo funciona com IA habilitada e continua operável quando o provedor falha.",
            ],
            [
                "7",
                "Executar testes ponta a ponta do CRM, corrigir usabilidade, desempenho percebido e mensagens de falha.",
                "Executar carga, isolamento, segurança, backup, restauração, métricas, alertas e runbooks.",
                "Release candidate sem bloqueadores P0 e demonstração ponta a ponta com dados representativos.",
            ],
            [
                "8",
                "Conduzir onboarding, treinamento curto, observação de uso e coleta de feedback dos pilotos.",
                "Preparar deploy, monitoramento, suporte técnico, limites de consumo e resposta a incidentes.",
                "Primeiro ciclo real de uma empresa piloto concluído e acompanhado pelos dois.",
            ],
            [
                "9",
                "Corrigir problemas P0 de CRM e experiência e medir ativação, uso semanal e economia de tempo.",
                "Corrigir problemas P0 de IA e plataforma e medir precisão, falhas, latência e custo por edital.",
                "Comparação entre linha de base e uso real, com causas das falhas e ações priorizadas.",
            ],
            [
                "10",
                "Consolidar onboarding comercial, materiais de uso, proposta inicial e melhorias de adoção.",
                "Consolidar capacidade, custos, SLOs, dívida técnica e plano de evolução da arquitetura.",
                "Decisão conjunta de cobrar, estender o piloto ou reposicionar, seguida do próximo roadmap.",
            ],
        ],
        widths=[1.1, 4.7, 4.7, 4.9],
        font_size=7.4,
    )

    add_heading(doc, "Distribuição de capacidade", 2)
    add_table(
        doc,
        ["Pessoa", "Foco principal", "Trabalho compartilhado", "Reserva operacional"],
        [
            ["Desenvolvedor A", "60 por cento CRM e experiência", "25 por cento contratos, testes e revisão", "15 por cento piloto, suporte e correções"],
            ["Desenvolvedor B", "60 por cento IA e plataforma", "25 por cento contratos, testes e revisão", "15 por cento operação, suporte e correções"],
        ],
        widths=[2.6, 4.2, 4.4, 4.2],
        font_size=8.3,
    )
    add_para(
        doc,
        "A reserva operacional não deve ser preenchida antecipadamente com novas funcionalidades. Ela absorve revisão, integração, falhas de ambiente, suporte ao piloto e correções necessárias para cumprir o resultado semanal.",
    )

    add_heading(doc, "Regras de dependência e integração", 2)
    add_bullets(doc, [
        "Na segunda feira, cada desenvolvedor assume um resultado principal e declara a dependência que precisa do outro.",
        "Contratos de API, eventos, schemas e migrações são combinados antes da implementação das duas pontas.",
        "Na quarta feira, o fluxo parcial é integrado; divergências não ficam acumuladas até a demonstração.",
        "Mudanças em tenant, autorização, entidades centrais e regras de preço exigem revisão do outro desenvolvedor.",
        "Bloqueio superior a meio dia gera trabalho conjunto temporário, com registro do impacto no objetivo semanal.",
        "Na sexta feira, ambos demonstram o resultado integrado, consultam métricas e escolhem no máximo duas correções prioritárias.",
    ])

    add_heading(doc, "Forma de trabalho", 2)
    add_bullets(doc, [
        "Planejamento semanal de 45 minutos com revisão de métricas e bloqueios.",
        "Demonstração interna no fim de cada semana usando dados reais.",
        "Pull request revisado pelo outro desenvolvedor para mudanças de domínio, segurança e dados.",
        "Decisões de arquitetura registradas em documentos curtos.",
        "Uma pessoa lidera a entrega e a outra conhece o caminho de operação e recuperação.",
        "Nenhuma frente permanece mais de duas semanas sem integração ponta a ponta.",
    ])

    add_heading(doc, "Definição de pronto", 2)
    add_checklist(doc, [
        "Regra de negócio e critério de aceite estão documentados.",
        "Implementação revisada.",
        "Testes relevantes passam.",
        "Migração e rollback foram considerados.",
        "Logs e métricas permitem diagnosticar falha.",
        "Isolamento por tenant foi verificado.",
        "Interface trata carregamento, vazio, erro e reprocessamento.",
        "Funcionalidade foi demonstrada com dado representativo.",
    ])

    add_heading(doc, "Ritual semanal sugerido", 2)
    add_table(
        doc,
        ["Momento", "Atividade"],
        [
            ["Segunda", "Escolher no máximo dois resultados da semana e dividir responsáveis"],
            ["Diário", "Atualização curta de progresso, bloqueio e próxima ação"],
            ["Quarta", "Integração intermediária e correção de divergências"],
            ["Sexta", "Demonstração, métricas, decisões e atualização do backlog"],
        ],
        widths=[3.2, 12.1],
        font_size=9,
    )

    # 15 Metrics
    add_heading(doc, "15 Métricas do produto e critérios de lançamento", 1)
    add_heading(doc, "Métrica principal", 2)
    add_para(
        doc,
        "A métrica principal sugerida é o número de oportunidades analisadas e revisadas até uma decisão comercial por tenant por mês. Ela mede uso completo, não apenas login, upload ou geração de texto.",
    )

    add_heading(doc, "Métricas por etapa", 2)
    add_table(
        doc,
        ["Etapa", "Métrica", "Pergunta respondida"],
        [
            ["Ativação", "Tempo até cadastrar catálogo e processar primeiro edital", "O onboarding entrega valor rápido"],
            ["Análise", "Tempo manual antes e depois do produto", "Existe economia mensurável"],
            ["IA", "Precisão por campo e Top 3 de matching", "A recomendação é útil"],
            ["Revisão", "Percentual de recomendações aceitas ou corrigidas", "Onde a IA falha"],
            ["CRM", "Oportunidades atualizadas semanalmente", "O CRM substitui controles paralelos"],
            ["Retenção", "Tenants ativos em quatro semanas consecutivas", "O valor se repete"],
            ["Operação", "Erros, p95, jobs falhos e tempo de recuperação", "O sistema é confiável"],
            ["Comercial", "Pilotos convertidos e disposição a pagar", "Existe negócio sustentável"],
        ],
        widths=[2.6, 6.0, 6.7],
        font_size=8.2,
    )

    add_heading(doc, "Critérios para iniciar o piloto", 2)
    add_checklist(doc, [
        "Escopo P0 concluído ou com exceções registradas.",
        "Empresa e catálogo configurados por tenant.",
        "Fluxo ponta a ponta demonstrado com edital real.",
        "IA atinge metas mínimas no conjunto de avaliação.",
        "Isolamento entre tenants aprovado.",
        "Backup restaurado com sucesso.",
        "Jobs falhos são visíveis e reprocessáveis.",
        "Existe responsável por suporte durante o piloto.",
    ])

    add_heading(doc, "Critérios para começar a cobrar", 2)
    add_checklist(doc, [
        "Pelo menos duas empresas concluíram um ciclo real.",
        "Usuários relatam economia de tempo e continuam usando sem cobrança diária do time.",
        "Falhas críticas possuem resposta operacional documentada.",
        "Termos, privacidade, limites de uso e suporte foram definidos.",
        "O custo de IA e infraestrutura por tenant é conhecido.",
        "Existe uma oferta simples com preço, limites e processo de onboarding.",
    ])

    # 16 Pilot
    add_heading(doc, "16 Estratégia de piloto e comercialização", 1)
    add_heading(doc, "Formato do piloto", 2)
    add_bullets(doc, [
        "Duração de quatro a seis semanas.",
        "Duas ou três empresas com perfis próximos.",
        "Um responsável nomeado em cada empresa.",
        "Onboarding acompanhado e coleta de linha de base.",
        "Reunião semanal curta para observar uso e impedimentos.",
        "Uso de editais e catálogos reais mediante acordo de confidencialidade quando necessário.",
        "Critérios de sucesso combinados antes do início.",
    ])

    add_heading(doc, "Dados que precisam ser coletados", 2)
    add_table(
        doc,
        ["Antes do produto", "Durante o piloto", "Depois do ciclo"],
        [
            ["Tempo médio de análise", "Tempo por etapa", "Tempo economizado"],
            ["Volume mensal de editais", "Editais processados", "Cobertura do fluxo"],
            ["Forma atual de controle", "Uso do CRM", "Controle paralelo mantido ou abandonado"],
            ["Erros mais frequentes", "Correções na IA", "Qualidade percebida"],
            ["Custo atual do processo", "Custo de IA", "Faixa de preço viável"],
        ],
        widths=[5.0, 5.1, 5.2],
        font_size=8.6,
    )

    add_heading(doc, "Oferta comercial inicial", 2)
    add_para(
        doc,
        "A primeira oferta deve ser simples. Uma assinatura por empresa pode incluir um número de usuários, editais processados e armazenamento, com excedente de IA claramente definido. "
        "O preço só deve ser fechado depois de medir custo por edital e valor percebido no piloto.",
    )

    add_heading(doc, "Perguntas para os clientes piloto", 2)
    add_numbered(doc, [
        "Qual etapa mais consome tempo atualmente.",
        "Qual erro custa mais caro.",
        "Quais dados precisam estar corretos para confiar na análise.",
        "Quem revisa o resultado da IA.",
        "Qual controle atual o sistema precisa substituir.",
        "Qual integração é realmente indispensável.",
        "Quanto tempo ou dinheiro precisa ser economizado para justificar a assinatura.",
    ])

    # 17 Risks
    page_break(doc)
    add_heading(doc, "17 Riscos e respostas", 1)
    add_table(
        doc,
        ["Risco", "Probabilidade", "Impacto", "Resposta"],
        [
            ["Escopo cresce durante o piloto", "Alta", "Alto", "Congelar P0 e trocar item quando uma exceção entrar"],
            ["Produto tenta atender qualquer setor", "Alta", "Alto", "Manter licitações como mercado inicial"],
            ["IA parece boa em poucos exemplos", "Alta", "Alto", "Dataset versionado e validação separada"],
            ["LLM inventa conclusão", "Média", "Alto", "Evidência obrigatória, confiança e revisão humana"],
            ["Vazamento entre tenants", "Baixa", "Crítico", "Testes de isolamento em banco, cache, jobs, vetores e arquivos"],
            ["Fila de IA prejudica CRM", "Média", "Alto", "Workers e filas separados com limites"],
            ["Dois desenvolvedores criam silos", "Média", "Médio", "Revisão cruzada e rotação de operação"],
            ["Dados fixos impedem onboarding", "Alta", "Alto", "Perfil empresarial e configuração por tenant"],
            ["CRM não substitui planilha", "Média", "Alto", "Observar uso, reduzir atrito e medir atualização semanal"],
            ["Custo de IA inviável", "Média", "Médio", "LLM apenas em casos ambíguos e custo por tenant"],
            ["Infraestrutura complexa demais", "Média", "Médio", "Manter monólito modular e serviços atuais"],
            ["Perda de dados", "Baixa", "Crítico", "Backup, retenção e restauração testada"],
        ],
        widths=[4.5, 2.2, 1.9, 6.8],
        font_size=7.7,
    )

    # 18 Backlog
    add_heading(doc, "18 Backlog priorizado", 1)
    add_heading(doc, "Prioridade zero", 2)
    add_checklist(doc, [
        "Congelar público, promessa e jornada do MVP.",
        "Criar perfil empresarial completo por tenant.",
        "Remover dados fixos de geração de documentos.",
        "Validar isolamento de tenant em banco, cache, arquivos, vetores e jobs.",
        "Consolidar atributos técnicos canônicos e tratamento de conflitos.",
        "Validar itens, kits, LPU e mínimo operacional com catálogo real.",
        "Fechar pipeline, descarte, reabertura, fases e encerramento.",
        "Separar claramente lista, detalhe e carregamento de abas do CRM.",
        "Criar dataset de avaliação da extração e do matching.",
        "Exigir evidência e confiança em cada recomendação técnica.",
        "Implementar feedback humano estruturado.",
        "Garantir fallback e feature flag para IA.",
        "Padronizar status, retries, idempotência e reprocessamento de jobs.",
        "Instrumentar métricas e alertas essenciais.",
        "Executar testes ponta a ponta, carga, isolamento e restauração.",
    ])

    add_heading(doc, "Prioridade um", 2)
    add_checklist(doc, [
        "Tarefas simples com responsável, prazo e status.",
        "Alertas críticos de prazo e pendência.",
        "Dashboard operacional enxuto.",
        "Chat por edital com citações.",
        "Importação de catálogo por planilha com relatório de erros.",
        "Modelos de checklist por tenant.",
        "Administração interna de tenants e consumo.",
        "Integração comercial solicitada por piloto confirmado.",
    ])

    add_heading(doc, "Prioridade dois", 2)
    add_checklist(doc, [
        "Automação de propostas e documentos avançados.",
        "Inteligência competitiva.",
        "Pós venda, contrato, entrega e faturamento.",
        "Customização de pipeline e campos.",
        "White label completo.",
        "Modelos especializados e fine tuning.",
        "Robô de lances com operação real.",
        "Arquitetura distribuída adicional quando houver métrica que a justifique.",
    ])

    # 19 Meeting
    add_heading(doc, "19 Pauta da reunião de alinhamento", 1)
    add_para(
        doc,
        "Tempo sugerido de reunião: duas horas. Se uma decisão não puder ser tomada, registrem responsável e prazo. Evitem transformar a reunião em revisão de telas ou código.",
    )
    add_table(
        doc,
        ["Tempo", "Tema", "Resultado esperado"],
        [
            ["10 minutos", "Objetivo comum", "Definir o que o produto precisa provar nos próximos 90 dias"],
            ["15 minutos", "Cliente inicial", "Escolher o perfil e listar empresas piloto"],
            ["20 minutos", "Jornada do MVP", "Aprovar o fluxo ponta a ponta"],
            ["20 minutos", "Corte de escopo", "Confirmar P0 e itens explicitamente adiados"],
            ["20 minutos", "CRM IA e Sistema", "Validar entregas e dependências das três frentes"],
            ["15 minutos", "Arquitetura", "Aprovar monólito modular, filas e isolamento"],
            ["10 minutos", "Responsabilidades", "Nomear liderança e revisão de cada frente"],
            ["10 minutos", "Datas e métricas", "Definir piloto, marcos e critérios de lançamento"],
        ],
        widths=[2.6, 4.4, 8.3],
        font_size=8.6,
    )

    add_heading(doc, "Perguntas que não podem ficar vagas", 2)
    add_checklist(doc, [
        "O produto é para qualquer empresa ou para qualquer empresa que participa de licitações.",
        "Qual problema o cliente pagará para resolver primeiro.",
        "Qual tela ou módulo atual pode ser ocultado sem prejudicar o MVP.",
        "Qual precisão mínima torna a IA útil e qual erro é inaceitável.",
        "Quem responde pelo suporte quando o piloto falhar.",
        "Qual custo mensal de infraestrutura e IA é aceitável.",
        "Qual compromisso de tempo cada desenvolvedor consegue manter.",
    ])

    add_heading(doc, "Encerramento da reunião", 2)
    add_checklist(doc, [
        "Backlog das próximas duas semanas selecionado.",
        "Responsável e revisor definidos para cada item.",
        "Empresa piloto e contato definidos.",
        "Data da primeira demonstração definida.",
        "Riscos sem resposta registrados com responsável.",
        "Próxima reunião agendada.",
    ])

    # 20 Decision log
    add_heading(doc, "20 Registro de decisões", 1)
    add_para(doc, "Preencher durante ou imediatamente após a reunião.")
    add_table(
        doc,
        ["Decisão", "Escolha", "Motivo", "Responsável", "Data"],
        [
            ["Cliente inicial", "", "", "", ""],
            ["Promessa do MVP", "", "", "", ""],
            ["Funcionalidades P0", "", "", "", ""],
            ["Itens fora do MVP", "", "", "", ""],
            ["Meta de qualidade da IA", "", "", "", ""],
            ["Provedor de IA inicial", "", "", "", ""],
            ["Divisão CRM IA Sistema", "", "", "", ""],
            ["Empresa piloto", "", "", "", ""],
            ["Data do piloto", "", "", "", ""],
            ["Critério para cobrança", "", "", "", ""],
        ],
        widths=[3.4, 3.1, 4.5, 2.6, 1.8],
        font_size=7.8,
    )

    add_heading(doc, "Dependências e impedimentos", 2)
    add_table(
        doc,
        ["Impedimento", "Impacto", "Próxima ação", "Responsável", "Prazo"],
        [["", "", "", "", ""] for _ in range(6)],
        widths=[3.8, 2.8, 4.3, 2.4, 1.9],
        font_size=8,
    )

    page_break(doc)

    # 21 30/60/90
    add_heading(doc, "21 Próximos passos em 30 60 e 90 dias", 1)
    add_heading(doc, "Primeiros 30 dias", 2)
    add_bullets(doc, [
        "Aprovar o escopo e congelar novas frentes.",
        "Selecionar empresas piloto.",
        "Completar perfil empresarial e onboarding por tenant.",
        "Validar catálogo, itens, kits e preços com dados reais.",
        "Definir e montar o dataset inicial de IA.",
        "Executar auditoria de isolamento e backup.",
    ])

    add_heading(doc, "Até 60 dias", 2)
    add_bullets(doc, [
        "Concluir o fluxo ponta a ponta.",
        "Atingir metas mínimas de extração e matching.",
        "Concluir testes de carga, segurança e recuperação.",
        "Criar suporte operacional e dashboards essenciais.",
        "Iniciar uso acompanhado com as empresas piloto.",
    ])

    add_heading(doc, "Até 90 dias", 2)
    add_bullets(doc, [
        "Medir economia de tempo, qualidade e retenção.",
        "Corrigir problemas P0 observados no piloto.",
        "Definir planos, limites e preço inicial.",
        "Converter pelo menos um piloto quando houver valor comprovado.",
        "Repriorizar o roadmap com base no uso, não em hipóteses internas.",
    ])

    add_heading(doc, "Resultado desejado ao final de 90 dias", 2)
    add_para(
        doc,
        "Ao final desse período, a equipe deve saber se empresas reais confiam na análise, usam o CRM, economizam tempo e aceitam pagar. "
        "Essa evidência orientará a próxima rodada de investimento técnico e comercial.",
    )

    # 22 References
    add_heading(doc, "22 Glossário e referências internas", 1)
    add_heading(doc, "Glossário", 2)
    add_table(
        doc,
        ["Termo", "Definição usada neste plano"],
        [
            ["Tenant", "Empresa cliente com dados e configurações isolados"],
            ["MVP", "Menor versão que prova valor com um fluxo completo e operável"],
            ["Matching", "Comparação entre requisito do edital e produto do catálogo"],
            ["RAG", "Recuperação de trechos relevantes antes da geração ou análise"],
            ["Evidência", "Trecho e origem que sustentam uma conclusão"],
            ["Job", "Processamento assíncrono com estado, progresso e tentativas"],
            ["Feature flag", "Configuração que habilita uma capacidade por ambiente ou tenant"],
            ["SLO", "Objetivo mensurável de qualidade ou desempenho do serviço"],
            ["RPO", "Quantidade máxima aceitável de dados perdidos em uma recuperação"],
            ["RTO", "Tempo máximo desejado para restaurar o serviço"],
        ],
        widths=[3.4, 11.9],
        font_size=8.8,
    )

    add_heading(doc, "Referências do repositório", 2)
    add_bullets(doc, [
        "readme.md com visão geral, stack, arquitetura e endpoints.",
        "docs plano_arquitetura_performance_crm_ia.md com diagnóstico de desempenho e arquitetura alvo.",
        "docs crm-design-system 07-backlog-imediato.md com pendências e regras atuais do CRM.",
        "docs crm-design-system 06-plano-execucao-redesign.md com decisões de produto e histórico do redesign.",
        "backend app market_profiles.py com perfis de licitações e matching documental.",
        "backend app auth models.py com tenants, usuários e papéis.",
        "backend app crm models.py com o modelo de domínio do CRM.",
        "docker-compose.prod.yaml com os componentes de produção e separação de workers.",
    ])

    add_heading(doc, "Conclusão", 2)
    add_para(
        doc,
        "A plataforma possui base suficiente para um MVP comercial. A prioridade dos próximos ciclos deve ser eliminar dependências fixas, fechar a jornada principal, medir a IA e provar o uso com empresas reais. "
        "A arquitetura atual suporta esse caminho desde que o time preserve o monólito modular, mantenha o processamento pesado fora da API e trate tenant, segurança e observabilidade como requisitos do produto.",
    )

    # Improve pagination and footer consistency
    for section in doc.sections:
        section.top_margin = Cm(2.0)
        section.bottom_margin = Cm(2.1)
        section.left_margin = Cm(2.2)
        section.right_margin = Cm(2.2)
        section.footer_distance = Cm(0.8)

    for paragraph in doc.paragraphs:
        if paragraph.style and paragraph.style.name.startswith("Heading"):
            paragraph.paragraph_format.keep_with_next = True

    return doc


if __name__ == "__main__":
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document = build_document()
    document.save(OUTPUT)
    print(OUTPUT)
