from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from create_mvp_improvement_plan_docx import (
    BLACK,
    MID_GRAY,
    add_bullets,
    add_checklist,
    add_heading,
    add_numbered,
    add_para,
    add_table,
    format_run,
    page_break,
    remove_paragraph_borders,
    style_document,
)


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_A = ROOT / "docs" / "Cronograma_Desenvolvedor_A_CRM_Experiencia.docx"
OUTPUT_B = ROOT / "docs" / "Cronograma_Desenvolvedor_B_IA_Plataforma.docx"


def add_cover(doc: Document, title: str, subtitle: str, purpose: str) -> None:
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    remove_paragraph_borders(p)
    run = p.add_run(title)
    format_run(run, bold=True, size=27, color=BLACK)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(20)
    run = p.add_run(subtitle)
    format_run(run, size=15, color=MID_GRAY)

    add_para(doc, "Plano individual vinculado ao MVP de CRM IA e Sistema", style="Small Note")
    add_para(doc, f"Versão preparada em {date.today().strftime('%d/%m/%Y')}", style="Small Note")
    doc.add_paragraph().paragraph_format.space_after = Pt(26)
    add_para(doc, purpose)
    add_para(
        doc,
        "O cronograma usa semanas como unidades de sequência. Se a dedicação for parcial, mantenha a ordem das entregas e amplie os prazos sem iniciar várias frentes incompletas.",
    )
    page_break(doc)


def add_common_operating_sections(doc: Document, partner_name: str, support_focus: str) -> None:
    add_heading(doc, "Rotina semanal", 1)
    add_table(
        doc,
        ["Momento", "Ação individual", "Interação com o outro desenvolvedor"],
        [
            ["Segunda feira", "Escolher um resultado principal, quebrar as tarefas e identificar o maior risco.", "Confirmar dependências, contratos e ordem de integração."],
            ["Diariamente", "Atualizar progresso, próxima ação, bloqueio e evidência produzida.", "Pedir ajuda quando um bloqueio ultrapassar meio dia."],
            ["Quarta feira", "Entregar uma versão parcial executável e revisar falhas.", "Integrar as duas frentes e corrigir divergências de contrato."],
            ["Sexta feira", "Demonstrar a entrega, consultar métricas e registrar pendências.", "Revisar o resultado do MVP e escolher no máximo duas correções."],
        ],
        widths=[2.7, 6.2, 6.4],
        font_size=8.5,
    )

    add_heading(doc, "Regras de colaboração", 2)
    add_bullets(doc, [
        f"{partner_name} revisa mudanças que afetem modelo de dados, contratos, tenant, segurança ou regras centrais.",
        "Schemas de API, eventos e migrações são combinados antes da implementação das duas pontas.",
        "Nenhuma integração crítica fica pendente por mais de uma semana.",
        "Um bloqueio superior a meio dia vira uma sessão conjunta temporária.",
        "A demonstração semanal usa dados representativos e inclui pelo menos um cenário de erro.",
        f"A reserva operacional cobre {support_focus}, revisão e correções; ela não recebe funcionalidades novas antecipadamente.",
    ])

    add_heading(doc, "Definição de pronto", 1)
    add_checklist(doc, [
        "Objetivo e critério de aceite estão documentados.",
        "Implementação foi revisada pelo outro desenvolvedor.",
        "Testes relevantes passam no ambiente de integração.",
        "Estados de sucesso, vazio, carregamento e falha foram tratados.",
        "Logs e métricas permitem diagnosticar o comportamento.",
        "Isolamento por tenant foi verificado quando a entrega manipula dados de negócio.",
        "Migração, compatibilidade e rollback foram considerados.",
        "Entrega foi demonstrada com dado representativo.",
        "Documentação operacional ou de uso foi atualizada quando necessária.",
    ])

    add_heading(doc, "Registro semanal", 1)
    add_para(doc, "Preencher no encerramento de cada semana para manter decisões, desvios e dependências visíveis.")
    add_table(
        doc,
        ["Campo", "Registro"],
        [
            ["Semana e período", ""],
            ["Resultado planejado", ""],
            ["Resultado entregue", ""],
            ["Evidência ou demonstração", ""],
            ["Métrica observada", ""],
            ["Bloqueios e causa", ""],
            ["Dependência do outro desenvolvedor", ""],
            ["Decisão tomada", ""],
            ["Pendência movida", ""],
            ["Objetivo da próxima semana", ""],
        ],
        widths=[5.0, 10.3],
        font_size=8.8,
    )


def build_developer_a() -> Document:
    doc = Document()
    style_document(doc)
    add_cover(
        doc,
        "Cronograma do Desenvolvedor A",
        "CRM experiência do usuário e validação do produto",
        "Este documento orienta a execução do Desenvolvedor A durante as dez semanas do MVP. O objetivo é entregar um CRM utilizável, integrar a análise de IA ao trabalho humano e comprovar que empresas piloto conseguem conduzir uma oportunidade até uma decisão.",
    )

    add_heading(doc, "1 Papel e resultado esperado", 1)
    add_para(
        doc,
        "O Desenvolvedor A lidera o CRM e a experiência. Sua responsabilidade termina quando o usuário consegue encontrar uma oportunidade, entender o que precisa fazer, revisar a recomendação, organizar itens e documentos e registrar o resultado sem depender de planilhas paralelas.",
    )
    add_heading(doc, "Responsabilidades permanentes", 2)
    add_bullets(doc, [
        "Jornada do usuário, navegação, estados de tela e consistência visual.",
        "Pipeline, oportunidade, fases, responsável, histórico, descarte, reabertura e encerramento.",
        "Catálogo, itens, kits, LPU, preço mínimo e documentos associados.",
        "Revisão humana da IA, evidências, confiança, correção e justificativa.",
        "Onboarding acompanhado, treinamento curto, coleta de feedback e métricas de adoção.",
        "Testes ponta a ponta do fluxo utilizado pelo cliente.",
    ])
    add_heading(doc, "Distribuição de capacidade", 2)
    add_table(
        doc,
        ["Uso da capacidade", "Percentual sugerido", "Aplicação"],
        [
            ["Entrega principal", "60 por cento", "CRM, experiência, catálogo e revisão humana"],
            ["Trabalho compartilhado", "25 por cento", "Contratos, modelo de dados, testes e revisão cruzada"],
            ["Reserva operacional", "15 por cento", "Piloto, suporte, integração e correções"],
        ],
        widths=[4.0, 3.3, 8.0],
        font_size=8.7,
    )

    add_heading(doc, "2 Cronograma de dez semanas", 1)
    add_table(
        doc,
        ["Semana", "Objetivo", "Atividades do Desenvolvedor A", "Dependência do Desenvolvedor B", "Critério de saída"],
        [
            ["1", "Congelar a jornada", "Mapear telas, fases, itens, kits, preços, documentos e encerramento. Identificar o que será ocultado do MVP.", "Arquitetura, entidades, limites de tenant e capacidades reais do pipeline de IA.", "Jornada aprovada, backlog P0 e critérios de aceite definidos."],
            ["2", "Ativar uma empresa", "Implementar perfil empresarial, preferências, onboarding e fluxo inicial de catálogo no frontend.", "Tenant, RBAC, auditoria, migrações e APIs de configuração estáveis.", "Empresa nova configura dados e acessa somente o próprio ambiente."],
            ["3", "Fechar a base do CRM", "Concluir lista, filtros, detalhe, fases, responsável, histórico e estados de carregamento, vazio e erro.", "Upload, parsing e jobs com status e erros legíveis para a interface.", "Edital enviado aparece no CRM e seu processamento pode ser acompanhado."],
            ["4", "Concluir itens e preços", "Finalizar vínculo item catálogo, kits, LPU, mínimo, descarte, reabertura e consistência entre lista e detalhe.", "Contratos de matching, candidatos, score, evidências e identificadores estáveis.", "Item real recebe candidatos sem perder quantidade, identidade ou preço."],
            ["5", "Entregar revisão humana", "Construir tela de recomendação, evidência, confiança, aceite, correção, justificativa e histórico da revisão.", "Dataset, métricas de qualidade e persistência do feedback estruturado.", "Correção humana fica registrada e pode alimentar avaliação posterior."],
            ["6", "Completar o trabalho operacional", "Concluir documentos, resultado por item, encerramento geral, tarefas e alertas P1 indispensáveis.", "Fallback, feature flags, limites, timeouts e sinalização clara da indisponibilidade da IA.", "Fluxo continua utilizável com IA disponível ou indisponível."],
            ["7", "Preparar a versão candidata", "Executar testes ponta a ponta, corrigir usabilidade, acessibilidade básica, desempenho percebido e mensagens de falha.", "Ambiente estável, dados de teste, observabilidade, segurança e restauração aprovadas.", "Demonstração completa sem bloqueadores P0."],
            ["8", "Iniciar o piloto", "Conduzir onboarding, treinamento, primeira sessão real, observação de uso e coleta de linha de base.", "Deploy, monitoramento, suporte técnico e resposta a incidentes disponíveis.", "Primeira empresa conclui um ciclo real acompanhado."],
            ["9", "Corrigir com evidência", "Resolver problemas P0 de CRM e experiência. Medir ativação, atualização semanal e economia de tempo.", "Relatório de falhas, latência e qualidade da IA para separar problema de interface e problema técnico.", "Melhorias priorizadas pela frequência e impacto observados no piloto."],
            ["10", "Preparar decisão comercial", "Consolidar onboarding, materiais de uso, proposta inicial, limites funcionais e melhorias de adoção.", "Custos, capacidade, SLOs, dívida técnica e limites operacionais consolidados.", "Decisão de cobrar, estender o piloto ou reposicionar registrada."],
        ],
        widths=[1.1, 2.7, 5.0, 3.7, 2.8],
        font_size=7.1,
    )

    add_heading(doc, "3 Entregas e métricas pessoais", 1)
    add_table(
        doc,
        ["Dimensão", "Indicador", "Meta inicial"],
        [
            ["Ativação", "Tempo até cadastrar catálogo e enviar o primeiro edital", "Medir e reduzir a cada piloto"],
            ["CRM", "Oportunidades atualizadas semanalmente", "Uso recorrente sem cobrança diária do time"],
            ["Revisão", "Recomendações aceitas ou corrigidas", "Toda correção possui motivo registrado"],
            ["Experiência", "Erros e bloqueios por etapa", "Nenhum bloqueador P0 aberto no piloto"],
            ["Desempenho", "Tempo até conteúdo útil e falhas de API percebidas", "Conteúdo principal em até 2,5 segundos em condição comum"],
            ["Valor", "Tempo economizado por oportunidade", "Comparação com a linha de base do cliente"],
        ],
        widths=[3.0, 6.1, 6.2],
        font_size=8.4,
    )
    add_heading(doc, "Checklist da primeira semana", 2)
    add_checklist(doc, [
        "Desenhar a jornada do login até o encerramento da oportunidade.",
        "Listar telas e abas que permanecem no MVP.",
        "Listar telas e módulos que serão ocultados ou adiados.",
        "Congelar fases, descarte, reabertura e resultados possíveis.",
        "Congelar regras de item, kit, preço mínimo e LPU.",
        "Definir estados de carregamento, vazio, falha e reprocessamento.",
        "Combinar contratos necessários com o Desenvolvedor B.",
        "Preparar uma demonstração do fluxo com dados representativos.",
    ])

    add_common_operating_sections(doc, "O Desenvolvedor B", "onboarding e suporte ao piloto")
    return doc


def build_developer_b() -> Document:
    doc = Document()
    style_document(doc)
    add_cover(
        doc,
        "Cronograma do Desenvolvedor B",
        "IA plataforma segurança e operação",
        "Este documento orienta a execução do Desenvolvedor B durante as dez semanas do MVP. O objetivo é tornar o processamento documental e o matching mensuráveis, seguros e operáveis, mantendo a arquitetura simples para uma equipe de dois desenvolvedores.",
    )

    add_heading(doc, "1 Papel e resultado esperado", 1)
    add_para(
        doc,
        "O Desenvolvedor B lidera IA e plataforma. Sua responsabilidade termina quando cada documento possui processamento rastreável, cada recomendação apresenta evidência e confiança, um tenant não acessa dados de outro e a equipe consegue diagnosticar, reprocessar e recuperar o serviço.",
    )
    add_heading(doc, "Responsabilidades permanentes", 2)
    add_bullets(doc, [
        "Upload, parsing, normalização, embeddings, matching e avaliação.",
        "Jobs, filas, retries, idempotência, reprocessamento e isolamento de workers.",
        "Gateway de IA, providers, timeout, fallback, feature flags, custo e limites.",
        "Tenant, RBAC, auditoria, arquivos privados, vetores e segurança.",
        "Métricas, logs, alertas, backup, restauração, runbooks e capacidade.",
        "Contratos de API e dados consumidos pelo CRM e pela revisão humana.",
    ])
    add_heading(doc, "Distribuição de capacidade", 2)
    add_table(
        doc,
        ["Uso da capacidade", "Percentual sugerido", "Aplicação"],
        [
            ["Entrega principal", "60 por cento", "IA, jobs, segurança, dados e plataforma"],
            ["Trabalho compartilhado", "25 por cento", "Contratos, modelo de dados, testes e revisão cruzada"],
            ["Reserva operacional", "15 por cento", "Operação, suporte, integração e correções"],
        ],
        widths=[4.0, 3.3, 8.0],
        font_size=8.7,
    )

    add_heading(doc, "2 Cronograma de dez semanas", 1)
    add_table(
        doc,
        ["Semana", "Objetivo", "Atividades do Desenvolvedor B", "Dependência do Desenvolvedor A", "Critério de saída"],
        [
            ["1", "Congelar a base técnica", "Mapear arquitetura, módulos, entidades, tenant, pipeline de IA, jobs, riscos, métricas e limites do MVP.", "Jornada, regras de CRM, dados exibidos e decisões que a IA precisa apoiar.", "Arquitetura e contratos principais aprovados pelos dois."],
            ["2", "Completar a base multiempresa", "Concluir tenant, RBAC, auditoria, migrações, configuração e remoção dos dados empresariais fixos.", "Campos de onboarding, preferências, catálogo e comportamento esperado da interface.", "Nova empresa é ativada sem código e passa nos testes de isolamento."],
            ["3", "Estabilizar documentos e jobs", "Concluir upload, hash, parsing, normalização, idempotência, retries, progresso, erro legível e reprocessamento.", "Estados de interface e contratos necessários para lista e detalhe do CRM.", "Edital processa de forma assíncrona e pode ser diagnosticado e reprocessado."],
            ["4", "Entregar matching base", "Implementar busca vetorial, candidatos, composição de score, evidência, versão e persistência do resultado.", "Regras de item, kit, preço, filtros e informações necessárias para decisão humana.", "Item real recebe candidatos com score e evidência rastreável."],
            ["5", "Tornar a IA mensurável", "Criar dataset versionado, separação de validação, métricas por campo, Top 3, falsos positivos, falsos negativos e ausência de evidência.", "Fluxo de aceite, correção e justificativa definido na revisão humana.", "Relatório reproduzível impede promoção de regressões acima do limite acordado."],
            ["6", "Controlar o uso de IA", "Concluir gateway, schema validado, fallback determinístico, feature flags, timeout, concorrência, custo e consumo por tenant.", "Mensagens e comportamento do CRM quando a IA estiver lenta ou indisponível.", "Sistema continua operável sem LLM e registra custo e falha de cada execução."],
            ["7", "Preparar a versão candidata", "Executar carga, isolamento, segurança, backup, restauração, métricas, alertas, runbooks e correções P0.", "Cenários ponta a ponta, dados de teste e comportamento esperado na interface.", "Release candidate sem bloqueadores P0 e recuperação testada."],
            ["8", "Operar o piloto", "Preparar deploy, dashboards, limites, suporte técnico, resposta a incidentes e acompanhamento de jobs reais.", "Agenda do piloto, empresa configurada, fluxo de onboarding e feedback dos usuários.", "Primeiro ciclo real termina com logs, custos e falhas observáveis."],
            ["9", "Corrigir com evidência", "Resolver problemas P0 de IA e plataforma. Medir precisão, falhas, latência, fila, reprocessamento e custo por edital.", "Erros percebidos pelo usuário, correções humanas e impacto no fluxo do CRM.", "Causas das falhas estão classificadas e ações são priorizadas por impacto."],
            ["10", "Preparar a próxima decisão", "Consolidar capacidade, custos, SLOs, segurança pendente, dívida técnica e evolução da arquitetura.", "Uso, retenção, economia de tempo, limites funcionais e proposta comercial.", "Decisão de cobrar, estender o piloto ou reposicionar registrada."],
        ],
        widths=[1.1, 2.7, 5.0, 3.7, 2.8],
        font_size=7.1,
    )

    add_heading(doc, "3 Entregas e métricas pessoais", 1)
    add_table(
        doc,
        ["Dimensão", "Indicador", "Meta inicial"],
        [
            ["Extração", "Campos críticos corretos", "95 por cento ou mais no conjunto revisado"],
            ["Matching", "Produto correto no Top 3", "85 por cento ou mais no conjunto revisado"],
            ["Evidência", "Conclusões técnicas com fonte", "100 por cento ou resultado não verificável"],
            ["Jobs", "Execuções com estado final explicável", "95 por cento ou mais"],
            ["API", "Latência p95 de operações comuns", "Abaixo de 500 milissegundos sem incluir jobs"],
            ["CRM", "Latência p95 da lista no servidor", "Abaixo de 800 milissegundos em base representativa"],
            ["Operação", "Restauração, alertas e diagnóstico", "Restauração testada e falhas críticas alertadas"],
            ["Custo", "Custo de IA por edital e tenant", "Conhecido antes de definir preço"],
        ],
        widths=[3.0, 6.1, 6.2],
        font_size=8.4,
    )
    add_heading(doc, "Checklist da primeira semana", 2)
    add_checklist(doc, [
        "Mapear componentes, módulos e dependências atuais.",
        "Congelar entidades e contratos necessários ao fluxo P0.",
        "Listar todas as tabelas, caches, arquivos, vetores e jobs que exigem tenant.",
        "Definir estados canônicos de job e política de reprocessamento.",
        "Definir versões de parser, embeddings, prompt e modelo.",
        "Definir métricas de extração, matching, latência, fila e custo.",
        "Combinar respostas de sucesso e erro necessárias ao Desenvolvedor A.",
        "Registrar riscos técnicos que podem bloquear o piloto.",
    ])

    add_common_operating_sections(doc, "O Desenvolvedor A", "operação e suporte ao piloto")
    return doc


def save_documents() -> None:
    OUTPUT_A.parent.mkdir(parents=True, exist_ok=True)
    doc_a = build_developer_a()
    doc_b = build_developer_b()
    doc_a.save(OUTPUT_A)
    doc_b.save(OUTPUT_B)
    print(OUTPUT_A)
    print(OUTPUT_B)


if __name__ == "__main__":
    save_documents()
