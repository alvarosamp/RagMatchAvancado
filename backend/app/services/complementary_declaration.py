"""Draft TOR declarations from saved match data; never updates the match."""
from io import BytesIO
from datetime import datetime
from zoneinfo import ZoneInfo
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from app.services import manual_match as manual


def catalog_for(item):
    return getattr(item, '_declaration_catalog', None) if hasattr(item, '_declaration_catalog') else getattr(item, 'catalog_product', None)


def meaningful(value):
    if isinstance(value, dict):
        return {k: clean for k, v in value.items() if (clean := meaningful(v)) not in (None, '', [], {})}
    if isinstance(value, list):
        return [clean for v in value if (clean := meaningful(v)) not in (None, '', [], {})]
    return value


def declaration_preview(notice, options):
    available = []
    for item in getattr(notice, 'notice_products', []) or []:
        catalog = catalog_for(item)
        if catalog:
            available.append({'id': item.id, 'item': item.item_number, 'sku': catalog.sku,
                              'brand': catalog.brand, 'model': catalog.model})
    ids = options.get('item_ids') or []
    if not isinstance(ids, list) or len(ids) != 1 or not isinstance(ids[0], str):
        return {'available_items': available, 'missing_fields': ['declaration.item_selection']}
    item = next((p for p in getattr(notice, 'notice_products', []) or [] if p.id == ids[0]), None)
    catalog = (catalog_for(item)) if item else None
    if catalog is None:
        return {'available_items': available, 'missing_fields': ['declaration.saved_match']}
    from app.services.match_item_export import build_crm_match_item_export
    state = (getattr(item, 'raw_payload', None) or {}).get('_manual_match')
    state = manual.validate_state(state) if state else manual.import_item(build_crm_match_item_export(notice, item), item)
    result = manual.comparison(state, manual.catalog_payload(catalog))
    points = []
    for req in state['requisitos']:
        detail = next((d for d in result['detalhes'] if d.get('id') == req['id']), {})
        status = detail.get('status', 'VERIFICAR')
        if status == 'NAO_ATENDE':
            treatment = 'Divergência técnica. Não declarar atendimento; avaliar produto alternativo ou nova evidência.'
        elif status == 'ATENDE':
            treatment = 'Atendimento registrado no match; conferir documento, versão e escopo antes da emissão.'
        else:
            treatment = '[A CONFIRMAR] Obter comprovação do fabricante ou validar a condição de fornecimento com a Tor.'
        points.append({'id': req['id'], 'text': req.get('texto_original') or '[A CONFIRMAR]',
                       'source': req.get('origem') or req.get('fonte') or req.get('evidencias') or '[A CONFIRMAR]',
                       'required': req.get('valor'), 'operator': req.get('operador'), 'condition': req.get('escopo') or req.get('condicao'), 'unit': req.get('unidade'),
                       'found': detail.get('encontrado'), 'status': status,
                       'reason': detail.get('motivo', 'Requisito pendente de revisão'), 'treatment': treatment})
    obligations = meaningful(state.get('obrigacoes_gerais') or {})
    return {'available_items': available, 'missing_fields': [] if points else ['declaration.requirements'],
            'item_id': item.id, 'item_number': item.item_number, 'item_description': item.description,
            'sku': catalog.sku, 'brand': catalog.brand, 'model': catalog.model,
            'catalog_revision': manual.catalog_payload(catalog).get('revisao'),
            'match_status': result['status'], 'requirements_hash': result['hash_requisitos'],
            'product_hash': result['hash_produto'], 'points': points, 'obligations': obligations,
            'documentary_pending': state.get('pendencias') or [], 'status': 'MINUTA_GERADA'}


def text(value):
    if value is None: return '[A CONFIRMAR]'
    if isinstance(value, bool): return 'Sim' if value else 'Não'
    if isinstance(value, dict): return '; '.join(f'{k}: {text(v)}' for k, v in value.items())
    if isinstance(value, list): return '; '.join(text(v) for v in value)
    return str(value)


def build_declaration(notice, company, options):
    context = declaration_preview(notice, options)
    if context['missing_fields']:
        raise ValueError('Salve um produto no match e selecione um único item com requisitos para gerar a minuta.')
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(3.3), Cm(2.8)
    section.left_margin = section.right_margin = Cm(2)
    normal = doc.styles['Normal'];normal.font.name = 'Arial';normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)
    for style in ('Title','Subtitle','Caption','Heading 1','Heading 2'):
        doc.styles[style].font.color.rgb = RGBColor(0,0,0)
    doc.styles['Title'].font.size = Pt(16)
    doc.styles['Subtitle'].font.size = Pt(10)
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
    doc.add_paragraph('Declaração complementar de especificações técnicas', 'Title')
    doc.add_paragraph('MINUTA PARA REVISÃO DA TOR', 'Subtitle')
    doc.add_paragraph('Documento gerado a partir da seleção salva no match. Os pontos pendentes exigem validação antes da emissão e assinatura. A geração desta minuta não altera o resultado técnico.')
    organ = getattr(getattr(notice,'organ',None),'name',None)
    bid = getattr(notice,'bid_number',None) or getattr(notice,'number',None)
    doc.add_paragraph(f"Órgão: {text(organ)}\nLicitação: {text(bid)} · Processo: {text(getattr(notice,'number',None))}\nItem: {text(context['item_number'])} · Fabricante: {text(context['brand'])}\nModelo: {text(context['model'])} · SKU: {text(context['sku'])}\nRevisão do produto: {text(context['catalog_revision'])}")
    doc.add_paragraph(f"Emissor proposto: {company.get('razao_social','[A CONFIRMAR]')}, CNPJ {company.get('cnpj','[A CONFIRMAR]')}, {company.get('endereco','[A CONFIRMAR]')}.")
    doc.add_paragraph('A Tor apresenta, para revisão de seu responsável, os pontos abaixo referentes ao equipamento identificado. As informações deverão ser conferidas nas fontes indicadas. Condições de fornecimento dependem de confirmação expressa; características técnicas pendentes devem ser comprovadas para este modelo e variante.')
    doc.add_paragraph('Conferência ponto a ponto', 'Heading 1')
    for point in context['points']:
        doc.add_paragraph(f"Requisito {point['id']}", 'Heading 2')
        table = doc.add_table(rows=0, cols=2)
        table.columns[0].width, table.columns[1].width = Cm(3.3), Cm(13.7)
        rows = [('Texto do edital',point['text']),('Fonte do requisito',text(point['source'])),('Valor exigido',(point['operator'] or '')+' '+text(point['required'])+(' '+point['unit'] if point['unit'] else '')),('Dado do produto',text(point['found'].get('valor',point['found']) if isinstance(point['found'],dict) else point['found'])),('Escopo e condição',text(point['condition'])),('Situação no match',point['status']+' - '+point['reason']),('Tratamento',point['treatment'])]
        for i,(label,value) in enumerate(rows):
            row=table.add_row();row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
            cells=row.cells;cells[0].text=label;cells[1].text=value
            for cell in cells:
                for paragraph in cell.paragraphs:paragraph.paragraph_format.keep_with_next=True
            for cell in cells:
                props=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
                for side in ('top','left','bottom','right'):
                    edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');borders.append(edge)
                props.append(borders)
            if i==0:
                for cell in cells:
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E8EEF4');cell._tc.get_or_add_tcPr().append(shade)
        doc.add_paragraph('Confirmação do responsável: [A CONFIRMAR]')
    if context['obligations']:
        doc.add_paragraph('Condições gerais de fornecimento', 'Heading 1');doc.add_paragraph(text(context['obligations']))
        doc.add_paragraph('Confirmação das condições pela Tor: [A CONFIRMAR]')
    if context['documentary_pending']:
        doc.add_paragraph('Pendências documentais', 'Heading 1');doc.add_paragraph(text(context['documentary_pending']))
    doc.add_paragraph('Validação e assinatura', 'Heading 1')
    doc.add_paragraph(f"Local e data: {options.get('city') or company.get('cidade') or '[A CONFIRMAR]'}, {options.get('date') or datetime.now(ZoneInfo('America/Sao_Paulo')).strftime('%d/%m/%Y')}.\nResponsável pela revisão: [A CONFIRMAR]\nRepresentante da Tor: {company.get('representante') or '[A CONFIRMAR]'}\nAssinatura: __________________________________")
    doc.add_paragraph('Rastreabilidade', 'Heading 1')
    doc.add_paragraph(f"Requisitos: {context['requirements_hash']}\nProduto: {context['product_hash']}\nEstado documental: MINUTA_GERADA",style='Caption')
    output=BytesIO();doc.save(output);return output.getvalue()
