"""Real HTTP validation of saved match selection and declaration scope."""
import io
import json
from docx import Document


def test_saved_review_candidate_generates_private_draft_without_approving_match(clients, unique):
    api=clients('empty','admin')
    def create(table,values):
        r=api.post(f'/crm/query/{table}/insert',json={'values':values});assert r.status_code==200,r.text
        return r.json()['data'][0]
    n=create('notices',{'number':'DECL-QA-'+unique,'title':'EXEMPLO SIMULADO - declaração'})
    p=create('catalog_products',{'sku':'DECL-'+unique,'brand':'EXEMPLO SIMULADO','model':'SW24','category':'switch','specification':json.dumps({'atributos':{'portas_rj45':{'valor':24}}})})
    i=create('notice_products',{'notice_id':n['id'],'item_number':'11','category':'switch','description':'Mínimo 24 portas RJ45','raw_payload':{'requisitos_atomicos':[{'id':'R1','campo_normalizado':'portas_rj45','operador':'>=','valor':24,'texto_original':'Mínimo 24 portas RJ45'}]}})
    base=f"/crm/notices/{n['id']}/products/{i['id']}/manual-match"
    loaded=api.get(base).json();body={'state':loaded['state'],'base_revision':loaded['base_revision'],'catalog_product_id':p['id']}
    preview=api.post(base+'/preview',json=body);assert preview.status_code==200,preview.text
    candidate=next(c for c in preview.json()['candidatos'] if c['id']==p['id']);body['catalog_hash']=candidate['hash_catalogo']
    saved=api.request('PUT',base,json=body);assert saved.status_code==200,saved.text
    changed=api.patch('/crm/query/catalog_products',json={'values':{'specification':json.dumps({'atributos':{'portas_rj45':{'valor':8}}})},'filters':[{'column':'id','op':'eq','value':p['id']}]});assert changed.status_code==200,changed.text
    req={'notice_id':n['id'],'template_id':'complementary_declaration','options':{'item_ids':[i['id']]},'idempotency_key':'DECL-'+unique}
    check=api.post('/documents/templates/complementary_declaration/preview',json=req);assert check.status_code==200,check.text
    ctx=check.json()['fields']['declaration'];assert ctx['sku']==p['sku'] and ctx['points'][0]['status']=='NAO_ATENDE'
    generated=api.post('/documents/generate',json=req);assert generated.status_code==200,generated.text
    doc=generated.json();assert doc['signature_status']=='not_requested'
    file=api.get(f"/documents/files/{doc['id']}/download");assert file.status_code==200,file.text
    word=Document(io.BytesIO(file.content));text=' '.join(x.text for x in word.paragraphs)+' '.join(c.text for t in word.tables for r in t.rows for c in r.cells)
    assert p['sku'] in text and 'Não declarar atendimento' in text and 'MINUTA' in text
    current=api.get(base).json();assert current['state']==saved.json()['state'] and current['base_revision']==saved.json()['base_revision']
    assert clients('alpha','admin').post('/documents/generate',json={**req,'idempotency_key':'FOREIGN-'+unique}).status_code==404
    req['options']={'item_ids':['foreign']};req['idempotency_key']='BAD-'+unique
    assert api.post('/documents/generate',json=req).status_code==400
    assert api.request('DELETE',f"/documents/files/{doc['id']}").status_code==204
