import csv
from io import BytesIO, StringIO
from types import SimpleNamespace

from openpyxl import load_workbook
from pypdf import PdfReader
import pytest

from app.services import export_service, analysis_export_service, tor_datasheet_generator


@pytest.fixture
def matching_report():
    return {"edital_id": 17, "total_requirements": 3, "results": [
        {"model": model, "score": score, "details": [{"attribute": "ports", "required": "48", "provided": "24", "status": status, "score": score, "reason": "QA evidência"}]}
        for model, score, status in [("QA-SW48", 1.0, "atende"), ("QA-SW24", 0.5, "verificar"), ("QA-SW8", 0, "nao_atende")]
    ]}


def test_excel_matching_report_can_be_opened_and_preserves_products(matching_report):
    workbook = load_workbook(BytesIO(export_service.export_xlsx(matching_report)))
    assert workbook.sheetnames == ["Resumo", "Detalhes"]
    values = [value for row in workbook["Resumo"].values for value in row]
    assert all(row["model"] in values for row in matching_report["results"])
    detail_values = [value for row in workbook["Detalhes"].values for value in row]
    assert "ports" in detail_values


def test_matching_pdf_has_readable_product_and_scores(matching_report):
    data = export_service.export_pdf(matching_report)
    text = "\n".join(page.extract_text() for page in PdfReader(BytesIO(data)).pages)
    assert "QA-SW48" in text and "QA-SW8" in text
    assert "100%" in text


def test_matching_csv_round_trips_all_statuses_and_products(matching_report):
    data = export_service.export_csv(matching_report)
    text = data.decode("utf-8-sig") if isinstance(data, bytes) else data
    rows = list(csv.reader(StringIO(text)))
    assert len(rows) >= 4
    assert all(item["model"] in text for item in matching_report["results"])


def analysis_document(empty=False):
    item = SimpleNamespace(id=1, item_number="1", numero_item="1", description="Switch QA 48 portas", descricao="Switch QA 48 portas",
        quantity=3, unit="UN", unit_price=100, total_value=300, has_risco=True, categoria="switch", municipio="São Paulo",
        uf="SP", data_disputa="2026-12-01", data_entrega="2026-12-10", is_exclusivo_me_epp=True,
        extracted_features={"ports": 48}, caracteristicas_bi={"quantidade_portas": "48"}, item_type="switch", raw={"riscos": ["Prazo curto"]}, item_payload={}, raw_payload={"quantidade_portas": "48"}, prazo_entrega="10 dias", unit_value=100)
    return SimpleNamespace(id=17, document_type="edital", source_filename="qa-edital.pdf", source_name="qa-edital.pdf", status="done", items=[] if empty else [item],
        created_at=None, updated_at=None, tenant_id=1, result={"edital": {"numero_pregao": "QA-17", "orgao": "Prefeitura QA", "riscos": ["Prazo curto"], "exclusivo_me_epp": True}, "documentos": ["Habilitação"], "itens_elegiveis": []})


@pytest.mark.parametrize("empty", [True, False])
def test_analysis_pdf_is_readable_for_empty_and_populated_documents(empty):
    document = analysis_document(empty)
    data = analysis_export_service.export_analysis_pdf(document)
    reader = PdfReader(BytesIO(data))
    assert len(reader.pages) >= 2
    text = "\n".join(page.extract_text() for page in reader.pages)
    assert "QA-17" in text
    if not empty:
        assert "Switch QA 48 portas" in text


@pytest.mark.parametrize("empty", [True, False])
def test_aggregate_analysis_pdf_has_period_and_document_inventory(empty):
    data = analysis_export_service.export_analysis_report_pdf([] if empty else [analysis_document()], period_label="QA outubro", generated_at="2026-10-10")
    text = "\n".join(page.extract_text() for page in PdfReader(BytesIO(data)).pages)
    assert "QA outubro" in text


@pytest.mark.parametrize("category,model,raw", [
    ("switch", "TOR-SW48", "48 portas Gigabit PoE 370 W IEEE 802.3af"),
    ("access_point", "TOR-AP6", "Wi-Fi 6 802.11ax 2.4 GHz e 5 GHz"),
    ("cable", "TOR-CAT6", "Cabo CAT6 UTP cobre 305 metros"),
    (None, None, "Equipamento genérico"),
])
def test_tor_datasheet_preview_and_real_pdf_preserve_identity(category, model, raw):
    extracted = {"model": model, "category": category, "manufacturer": "Fabricante QA", "raw_text": raw,
                 "specs": {"ports": 48, "speed": "1 Gbps", "poe": True, "warranty": "36 meses"}}
    preview = tor_datasheet_generator.build_tor_datasheet_preview(extracted, pn_tor=model, category=category)
    assert preview["fonte"]["fabricante"] == "Fabricante QA"
    if model:
        assert preview["pn_tor"] == model
    data = tor_datasheet_generator.export_tor_datasheet_pdf(preview)
    text = "\n".join(page.extract_text() for page in PdfReader(BytesIO(data)).pages)
    assert "TOR" in text and "Datasheet" in text
